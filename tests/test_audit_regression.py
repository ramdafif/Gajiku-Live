import ast
import csv
import io
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["START_INLINE_WORKER"] = "0"
os.environ["EXPORT_CSV_ENABLED"] = "0"
from app import create_app
from app.db import get_db
from app.routes import web
from app.services import email_service
from app.services.security_service import business_today, token_serializer, credential_version
from app.utils.cache import _cache
from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash, check_password_hash

PASSWORD = "FixturePassword123"
HASH = generate_password_hash(PASSWORD)
TODAY = business_today().isoformat()
PERIOD = TODAY[:7]
CONFIG = {"TESTING": True, "DB_TYPE": "sqlite", "APP_SECRET": "isolated-fixture-secret-not-for-production",
          "SESSION_COOKIE_SECURE": False, "PUBLIC_BASE_URL": "https://fixture.example",
          "ADMIN_PASSWORD": "", "SUPERADMIN_PASSWORD": "", "SEED_DEMO_DATA": False,
          "ALLOW_DATE_OVERRIDE": False, "ALLOW_DESTRUCTIVE_RESET": False, "ENABLE_TEST_ENDPOINTS": False}

class AuditRegression(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix="gajiku-audit-")
        cls.base = Path(cls.directory.name) / "base.db"
        app = create_app({**CONFIG, "DB_PATH": str(cls.base)})
        with app.app_context():
            db = get_db()
            for name in ("Fixture PT One", "Fixture PT Two"):
                db.execute("INSERT INTO clients (name,name_key,allow_projects,created_at,updated_at) VALUES (?,?,1,?,?)", (name,name.lower(),TODAY,TODAY))
            for client_id in (1,2):
                db.execute("INSERT INTO client_projects (client_id,name,name_key,created_at,updated_at) VALUES (?,?,?,?,?)", (client_id,f"Project {client_id}",f"project {client_id}",TODAY,TODAY))
            for admin_id,name,email,company,role,active in [(1,"Fixture Super","super@fixture.test",None,"superadmin",1),(2,"Fixture Admin One","admin@example.com","Fixture PT One","admin",1),(3,"Fixture Admin Two","admin2@fixture.test","Fixture PT Two","admin",1),(4,"Fixture Disabled","disabled@fixture.test","Fixture PT One","admin",0),(5,"Fixture Unassigned","empty@fixture.test",None,"admin",1)]:
                db.execute("INSERT INTO admins (id,name,email,password_hash,company,role,status_aktif,hak_approval,created_at) VALUES (?,?,?,?,?,?,?,1,?)",(admin_id,name,email,HASH,company,role,active,TODAY))
            for employee_id in (1,2,3):
                company = f"Fixture PT {'One' if employee_id != 2 else 'Two'}"
                project = f"Project {1 if employee_id != 2 else 2}"
                email = f"employee{employee_id}@fixture.test"
                db.execute("INSERT INTO pegawai (id,id_pegawai,nama,email,jabatan,gaji,status_aktif,perusahaan,perusahaan_induk,no_rekening,no_telp,siklus_gaji,created_at) VALUES (?,?,?,?,?,6000000,1,?,?,?,'0800','A',?)",(employee_id,f"E{employee_id}",f"Employee Fixture {employee_id}",email,"Fixture",project,company,"1234",TODAY))
                if employee_id != 3:
                    db.execute("INSERT INTO user_accounts (pegawai_id,name,email,password_hash,status_aktif,created_at) VALUES (?,?,?,?,1,?)",(employee_id,f"Employee Fixture {employee_id}",email,HASH,TODAY))
                    db.execute("INSERT INTO users (id,name,email,gaji,created_at) VALUES (?,?,?,6000000,?)",(employee_id,f"Employee Fixture {employee_id}",email,TODAY))
                    for offset,product in ((0,"reg"),(10,"urg")):
                        db.execute("INSERT INTO transactions (id,user_id,tanggal,periode,nominal,admin_fee,status,created_at,product,cancel_until) VALUES (?,?,?,?,10000,15000,'on-proses',?,?,?)",(employee_id+offset,employee_id,TODAY,PERIOD,TODAY+'T09:00:00',product,(datetime.now()+timedelta(seconds=25)).isoformat()))
            db.commit()
    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()
    def setUp(self):
        _cache.clear()
        self.path = Path(self.directory.name) / (self._testMethodName + ".db")
        shutil.copyfile(self.base,self.path)
        self.app = create_app({**CONFIG,"DB_PATH":str(self.path)})
        self.client = self.app.test_client()
    def query(self, sql, params=()):
        with self.app.app_context():
            db = get_db()
            cursor = db.execute(sql,params)
            rows = cursor.fetchall() if cursor.description else []
            db.commit()
            return [dict(row) for row in rows]
    def csrf(self, path="/login", client=None):
        response = (client or self.client).get(path)
        match = re.search(r'name="csrf_token" value="([^"]+)"', response.get_data(as_text=True))
        self.assertIsNotNone(match, (path,response.status_code))
        return match.group(1)
    def login(self, email="admin@example.com", client=None):
        target = client or self.client
        response = target.post("/login",data={"email":email,"password":PASSWORD,"csrf_token":self.csrf(client=target)})
        self.assertEqual(response.status_code,302)
        target.get(response.location)
        return response
    def post(self,path,data=None,client=None):
        target = client or self.client
        with target.session_transaction() as session:
            token = session.get("_csrf_token")
        return target.post(path,data={**(data or {}),"csrf_token":token})
    def test_all_admin_get_pages_and_warm_cache(self):
        self.login()
        for repeat in range(2):
            for path in ("/admin/dashboard","/admin/settings","/admin/pegawai","/admin/riwayat","/admin/export","/admin/pegawai/import/template","/admin/projects","/health","/metrics"):
                response=self.client.get(path)
                self.assertEqual(response.status_code,200,(path,repeat))
        html=self.client.get('/admin/dashboard').get_data(as_text=True)
        self.assertLess(html.index('Permintaan Reg On-Proses'),html.index('Total Pegawai (master)'))
        self.assertLess(html.index('Permintaan Urg On-Proses'),html.index('Total Pegawai (master)'))
        self.assertEqual(self.client.get('/').location,'/admin/dashboard')
    def test_admin_tenant_isolation_all_reads(self):
        self.login()
        for path in ('/admin/dashboard','/admin/riwayat','/admin/export','/admin/pegawai'):
            html=self.client.get(path).get_data(as_text=True)
            self.assertIn('Employee Fixture 1',html,path)
            self.assertNotIn('Employee Fixture 2',html,path)
        self.assertEqual(self.client.get('/superadmin/clients').status_code,403)
    def test_no_company_admin_does_not_crash_or_leak(self):
        self.login('empty@fixture.test')
        for repeat in range(2):
            for path in ('/admin/dashboard','/admin/riwayat','/admin/export','/admin/pegawai'):
                response=self.client.get(path)
                self.assertEqual(response.status_code,200,path)
                self.assertNotIn('Employee Fixture 1',response.get_data(as_text=True),path)
    def test_superadmin_pages_and_cache(self):
        self.login('super@fixture.test')
        for repeat in range(2):
            for path in ('/superadmin/dashboard','/superadmin/settings','/superadmin/admins','/superadmin/clients','/superadmin/clients/1','/superadmin/riwayat','/admin/pegawai','/admin/dashboard','/superadmin/export_range?start='+TODAY+'&end='+TODAY):
                self.assertEqual(self.client.get(path).status_code,200,(path,repeat))
    def test_inactive_admin_login_and_existing_session_revoked(self):
        response=self.client.post('/login',data={'email':'disabled@fixture.test','password':PASSWORD,'csrf_token':self.csrf()})
        self.assertEqual(response.status_code,200)
        with self.client.session_transaction() as session:
            self.assertFalse(session.get('is_admin'))
        self.login()
        self.query('UPDATE admins SET status_aktif=0 WHERE id=2')
        self.assertEqual(self.client.get('/admin/dashboard').location,'/login')
    def test_admin_scope_and_permission_updates_take_effect(self):
        self.login()
        self.query('UPDATE admins SET company=?,hak_approval=0 WHERE id=2',('Fixture PT Two',))
        html=self.client.get('/admin/dashboard').get_data(as_text=True)
        self.assertIn('Employee Fixture 2',html)
        self.assertNotIn('Employee Fixture 1',html)
        self.assertNotIn('>Approve</button>',html)
        self.post('/admin/tx/2/approve')
        self.assertIsNone(self.query('SELECT admin_approved_at FROM transactions WHERE id=2')[0]['admin_approved_at'])
    def test_destructive_and_global_settings_denied(self):
        self.login()
        before=self.query('SELECT COUNT(*) AS total FROM transactions')[0]['total']
        for path in ('/reset','/admin/reset','/admin/reset_all','/superadmin/reset','/superadmin/reset_all'):
            self.assertEqual(self.post(path).status_code,403,path)
        self.post('/admin/products',{'products':'reg'})
        self.assertEqual(self.query("SELECT value FROM app_settings WHERE key='enabled_products'")[0]['value'],'reg,urg')
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM transactions')[0]['total'],before)
        self.assertEqual(self.post('/txn/test').status_code,404)
    def test_csrf_login_and_mutations(self):
        self.assertEqual(self.client.post('/login',data={'email':'admin@example.com','password':PASSWORD}).status_code,400)
        self.assertEqual(self.client.post('/register',data={}).status_code,400)
        self.assertEqual(self.client.post('/forgot',data={}).status_code,400)
        self.login()
        self.assertEqual(self.client.post('/admin/tx/1/approve').status_code,400)
    def test_employee_pages_and_status_idor(self):
        self.login('employee1@fixture.test')
        for path in ('/dashboard','/settings','/tarik-gaji','/riwayat'):
            self.assertEqual(self.client.get(path).status_code,200,path)
        response=self.client.get('/api/tx_status?ids=1,2,11,12')
        self.assertEqual({row['id'] for row in response.json['items']},{1,11})
        self.assertEqual(self.client.get('/api/tx_status?ids='+','.join(str(value) for value in range(101))).status_code,400)
    def test_cancel_deadline_ownership_and_approval(self):
        self.login('employee1@fixture.test')
        self.query('UPDATE transactions SET cancel_until=? WHERE id=1',((datetime.now()-timedelta(seconds=1)).isoformat(),))
        self.post('/tx/1/cancel')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=1')[0]['status'],'on-proses')
        self.query('UPDATE transactions SET cancel_until=?,admin_approved_at=? WHERE id=1',((datetime.now()+timedelta(seconds=30)).isoformat(),TODAY))
        self.post('/tx/1/cancel')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=1')[0]['status'],'on-proses')
        self.post('/tx/2/cancel')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=2')[0]['status'],'on-proses')
        self.query('UPDATE transactions SET admin_approved_at=NULL WHERE id=1')
        self.post('/tx/1/cancel')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=1')[0]['status'],'dibatalkan')
    def test_approval_layers_and_repeated_final(self):
        self.login('super@fixture.test')
        self.post('/superadmin/tx/1/approve')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=1')[0]['status'],'on-proses')
        self.post('/superadmin/logout')
        self.login()
        self.post('/admin/tx/2/approve')
        self.assertIsNone(self.query('SELECT admin_approved_at FROM transactions WHERE id=2')[0]['admin_approved_at'])
        self.post('/admin/tx/1/approve')
        self.assertIsNotNone(self.query('SELECT admin_approved_at FROM transactions WHERE id=1')[0]['admin_approved_at'])
        self.post('/admin/tx/1/reject')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=1')[0]['status'],'on-proses')
        self.post('/admin/logout')
        self.login('super@fixture.test')
        self.post('/superadmin/tx/1/approve')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=1')[0]['status'],'sukses')
        timestamp=self.query('SELECT final_approved_at FROM transactions WHERE id=1')[0]['final_approved_at']
        self.post('/superadmin/tx/1/approve')
        self.post('/superadmin/tx/1/reject')
        self.assertEqual(self.query('SELECT final_approved_at,status FROM transactions WHERE id=1')[0],{'final_approved_at':timestamp,'status':'sukses'})
    def withdrawal_data(self,client=None):
        html=(client or self.client).get('/tarik-gaji').get_data(as_text=True)
        token=re.search(r'name="request_token" value="([^"]+)"',html).group(1)
        return {'request_token':token,'produk':'reg','nominal':'100.000','rekening_tujuan_key':'bank_utama','tanggal':'31'}
    def test_withdrawal_date_amount_and_idempotency(self):
        self.login('employee1@fixture.test')
        data=self.withdrawal_data()
        before=self.query('SELECT COUNT(*) AS total FROM transactions')[0]['total']
        with patch.object(web,'enqueue_email'):
            self.post('/tarik-gaji',data)
            self.post('/tarik-gaji',data)
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM transactions')[0]['total'],before+1)
        self.assertEqual(self.query('SELECT tanggal FROM transactions WHERE request_key IS NOT NULL')[0]['tanggal'],TODAY)
        data=self.withdrawal_data()
        data['nominal']='-100000'
        self.post('/tarik-gaji',data)
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM transactions')[0]['total'],before+1)
    def test_concurrent_withdrawal_limit(self):
        self.query('DELETE FROM transactions')
        clients=[self.app.test_client(),self.app.test_client()]
        requests=[]
        for client in clients:
            self.login('employee1@fixture.test',client)
            data=self.withdrawal_data(client)
            data.update(produk='urg',nominal='3000000')
            requests.append(data)
        barrier=threading.Barrier(2)
        def submit(index):
            barrier.wait()
            return self.post('/tarik-gaji',requests[index],clients[index]).status_code
        with patch.object(web,'enqueue_email'), ThreadPoolExecutor(max_workers=2) as pool:
            statuses=list(pool.map(submit,range(2)))
        self.assertTrue(all(status in (200,302) for status in statuses),statuses)
        self.assertEqual(self.query('SELECT COUNT(*) AS total,SUM(nominal) AS amount FROM transactions')[0],{'total':1,'amount':3000000})
    def test_account_and_master_inactive_blocks_withdrawal(self):
        self.login('employee1@fixture.test')
        data=self.withdrawal_data()
        self.query('UPDATE pegawai SET status_aktif=0 WHERE id=1')
        self.assertEqual(self.post('/tarik-gaji',data).status_code,302)
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM transactions WHERE request_key IS NOT NULL')[0]['total'],0)
    def test_forgot_password_requires_email_proof_and_single_use(self):
        old=self.query('SELECT password_hash FROM user_accounts WHERE pegawai_id=1')[0]['password_hash']
        with patch.object(email_service,'send_email',return_value=True) as send:
            self.client.post('/forgot',data={'email':'employee1@fixture.test','csrf_token':self.csrf('/forgot')})
            body=send.call_args.args[1]
        self.assertEqual(self.query('SELECT password_hash FROM user_accounts WHERE pegawai_id=1')[0]['password_hash'],old)
        url=re.search(r'https://fixture.example(/reset-password/[^\s]+)',body).group(1)
        csrf=self.csrf(url)
        response=self.client.post(url,data={'csrf_token':csrf,'new_password':'NewFixture123','new_password2':'NewFixture123'})
        self.assertEqual(response.status_code,302)
        self.assertTrue(check_password_hash(self.query('SELECT password_hash FROM user_accounts WHERE pegawai_id=1')[0]['password_hash'],'NewFixture123'))
        self.assertEqual(self.client.get(url).location,'/forgot')
    def test_registration_requires_email_verification(self):
        with patch.object(email_service,'send_email',return_value=True) as send:
            response=self.client.post('/register',data={'csrf_token':self.csrf('/register'),'name':'New Fixture','email':'employee3@fixture.test','password':PASSWORD,'password2':PASSWORD})
            self.assertEqual(response.status_code,302)
            body=send.call_args.args[1]
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM user_accounts WHERE pegawai_id=3')[0]['total'],0)
        url=re.search(r'https://fixture.example(/verify-registration/[^\s]+)',body).group(1)
        token=self.csrf(url)
        self.client.post(url,data={'csrf_token':token})
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM user_accounts WHERE pegawai_id=3')[0]['total'],1)
        self.assertEqual(self.client.get(url).location,'/register')
    def test_rate_limit_is_persistent(self):
        csrf=self.csrf()
        for index in range(20):
            self.assertNotEqual(self.client.post('/login',data={'csrf_token':csrf,'email':'unknown@fixture.test','password':'wrong'}).status_code,429)
        self.assertEqual(self.client.post('/login',data={'csrf_token':csrf,'email':'unknown@fixture.test','password':'wrong'}).status_code,429)
    def test_csv_formula_injection(self):
        self.query("UPDATE users SET name='=1+1' WHERE id=1")
        self.login()
        rows=list(csv.reader(io.StringIO(self.client.get('/admin/export').get_data(as_text=True).lstrip('\ufeff'))))
        self.assertTrue(any("'=1+1" in row for row in rows))
    def test_employee_financial_history_cannot_be_deleted(self):
        self.login()
        self.post('/admin/pegawai/1/delete')
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM pegawai WHERE id=1')[0]['total'],1)
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM transactions WHERE user_id=1')[0]['total'],2)
    def test_settings_password_change_revokes_other_sessions(self):
        second=self.app.test_client()
        self.login(client=second)
        self.login()
        self.csrf('/admin/settings')
        self.post('/admin/settings',{'old_password':PASSWORD,'new_password':'ChangedFixture123','new_password2':'ChangedFixture123'})
        self.assertEqual(self.client.get('/admin/dashboard').status_code,200)
        self.assertEqual(second.get('/admin/dashboard').location,'/login')
    def test_env_super_settings_does_not_edit_first_admin(self):
        self.app.config.update(SUPERADMIN_USERNAME='envsuper@fixture.test',SUPERADMIN_PASSWORD=PASSWORD)
        self.login('envsuper@fixture.test')
        old=self.query('SELECT password_hash FROM admins WHERE id=1')[0]['password_hash']
        self.assertIn('Akun Superadmin ENV',self.client.get('/superadmin/settings').get_data(as_text=True))
        self.post('/superadmin/settings',{'old_password':PASSWORD,'new_password':'OtherFixture123','new_password2':'OtherFixture123'})
        self.assertEqual(self.query('SELECT password_hash FROM admins WHERE id=1')[0]['password_hash'],old)
    def test_ppn_products_and_admin_status_create(self):
        self.login('super@fixture.test')
        self.post('/superadmin/settings',{'form_type':'ppn','ppn_enabled':'0'})
        self.assertEqual(self.query("SELECT value FROM app_settings WHERE key='ppn_enabled'")[0]['value'],'0')
        self.post('/superadmin/admins/add',{'nama':'New Admin','email':'newadmin@fixture.test','password':PASSWORD,'perusahaan':'Fixture PT One','status':'0','cutoff_hari':'2','produk_reg_aktif':'1'})
        self.assertEqual(self.query("SELECT status_aktif FROM admins WHERE email='newadmin@fixture.test'")[0]['status_aktif'],0)
    def test_clients_projects_scope_and_used_project_protection(self):
        self.login()
        self.post('/admin/projects/2/delete')
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM client_projects WHERE id=2')[0]['total'],1)
        self.post('/admin/projects/1/delete')
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM client_projects WHERE id=1')[0]['total'],1)
        self.post('/admin/projects/add',{'name':'Fleminh'})
        self.assertEqual(self.query("SELECT name FROM client_projects WHERE client_id=1 AND name_key='fleming'")[0]['name'],'Fleming')
    def test_avatar_validation_and_internal_redirect(self):
        self.login()
        response=self.post('/avatar/upload',{'avatar':(io.BytesIO(b'not-an-image'),'payload.png')})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.query('SELECT avatar_path FROM admins WHERE id=2')[0]['avatar_path'],'')
    def test_template_excel_and_import(self):
        self.login()
        response=self.client.get('/admin/pegawai/import/template')
        workbook=load_workbook(io.BytesIO(response.data))
        self.assertEqual(workbook.active.max_column,14)
        workbook.close()
        workbook=Workbook()
        workbook.active.append(['ID','Nama','Email','Jabatan','Project','Rek','Rek2','EW','Telp','Gaji','Status','Siklus','Fee','Perusahaan'])
        workbook.active.append(['NEW1','Imported Fixture','new@fixture.test','Staff','Project 1','1234','','','0800',4000000,'Aktif','A',15000,'Fixture PT One'])
        output=io.BytesIO()
        workbook.save(output)
        output.seek(0)
        self.post('/admin/pegawai/import',{'file':(output,'fixture.xlsx')})
        self.assertEqual(self.query("SELECT perusahaan_induk FROM pegawai WHERE id_pegawai='NEW1'")[0]['perusahaan_induk'],'Fixture PT One')
    def test_windu_cannot_bypass_admin_layer(self):
        self.query("UPDATE pegawai SET perusahaan_induk='Windu Karya' WHERE id=1")
        self.query("UPDATE admins SET company='Windu Karya',hak_approval=0 WHERE id=2")
        self.login('super@fixture.test')
        self.post('/superadmin/tx/1/approve')
        self.assertEqual(self.query('SELECT status FROM transactions WHERE id=1')[0]['status'],'on-proses')
    def test_cross_company_employee_edit_denied(self):
        self.login()
        self.post('/admin/pegawai/2/update',{'id_pegawai':'E2','nama':'Changed','email':'employee2@fixture.test','gaji':'1','status':'1','perusahaan_induk':'Fixture PT One','perusahaan':'Project 1'})
        self.assertEqual(self.query('SELECT gaji FROM pegawai WHERE id=2')[0]['gaji'],6000000)
    def test_legacy_blanks_preserved_on_edit(self):
        self.query("UPDATE pegawai SET perusahaan_induk='',perusahaan='' WHERE id=3")
        self.login('super@fixture.test')
        self.post('/admin/pegawai/3/update',{'id_pegawai':'E3','nama':'Updated Fixture','email':'employee3@fixture.test','gaji':'5000000','status':'1','siklus_gaji':'A','perusahaan_induk':'','perusahaan':''})
        row=self.query('SELECT nama,perusahaan_induk,perusahaan FROM pegawai WHERE id=3')[0]
        self.assertEqual(row,{'nama':'Updated Fixture','perusahaan_induk':'','perusahaan':''})
    def test_client_mou_and_disallow_validation(self):
        self.login('super@fixture.test')
        self.post('/superadmin/clients',{'name':'Invalid MoU','mou_start':'2026-12-31','mou_end':'2026-01-01'})
        self.assertEqual(self.query("SELECT COUNT(*) AS total FROM clients WHERE name='Invalid MoU'")[0]['total'],0)
        self.post('/superadmin/clients/1',{'name':'Fixture PT One','allow_projects':'0'})
        self.assertEqual(self.query('SELECT allow_projects FROM clients WHERE id=1')[0]['allow_projects'],1)
    def test_inactive_admin_env_cannot_bypass(self):
        self.app.config.update(ADMIN_EMAIL='disabled@fixture.test',ADMIN_PASSWORD=PASSWORD)
        response=self.client.post('/login',data={'csrf_token':self.csrf(),'email':'disabled@fixture.test','password':PASSWORD})
        self.assertEqual(response.status_code,200)
        with self.client.session_transaction() as session:
            self.assertFalse(session.get('is_admin'))
    def test_valid_avatar_is_reencoded_to_temporary_directory(self):
        from PIL import Image
        self.app.static_folder=str(Path(self.directory.name)/'static')
        self.login()
        image=io.BytesIO()
        Image.new('RGB',(12,12),'red').save(image,format='PNG')
        image.seek(0)
        self.post('/avatar/upload',{'avatar':(image,'avatar.png')})
        path=self.query('SELECT avatar_path FROM admins WHERE id=2')[0]['avatar_path']
        self.assertTrue((Path(self.app.static_folder)/path).exists())
        self.assertIn(path,self.client.get('/admin/dashboard').get_data(as_text=True))
    def test_approved_transaction_cannot_be_cancelled_in_polling(self):
        self.login('employee1@fixture.test')
        self.query('UPDATE transactions SET admin_approved_at=? WHERE id=1',(TODAY,))
        item=self.client.get('/api/tx_status?ids=1').json['items'][0]
        self.assertFalse(item['can_cancel'])
        self.assertEqual(item['remaining'],0)
    def test_superadmin_password_does_not_accept_env_admin_password(self):
        self.app.config['ADMIN_PASSWORD']='WrongEnvPassword'
        self.login('super@fixture.test')
        self.post('/superadmin/settings',{'old_password':'WrongEnvPassword','new_password':'ShouldNotChange123','new_password2':'ShouldNotChange123'})
        self.assertTrue(check_password_hash(self.query('SELECT password_hash FROM admins WHERE id=1')[0]['password_hash'],PASSWORD))
    def test_login_sql_injection_is_not_authentication(self):
        response=self.client.post('/login',data={'csrf_token':self.csrf(),'email':"' OR 1=1 --",'password':'invalid'})
        self.assertEqual(response.status_code,200)
        with self.client.session_transaction() as session:
            self.assertFalse(session.get('is_admin'))
    def test_duplicate_employee_returns_conflict_not_internal_error(self):
        from app.db import get_db
        @self.app.post('/fixture-duplicate')
        def duplicate():
            get_db().execute('INSERT INTO clients (name,name_key,created_at,updated_at) VALUES (?,?,?,?)',('Duplicate','fixture pt one',TODAY,TODAY))
            return 'unexpected'
        self.csrf()
        self.assertEqual(self.post('/fixture-duplicate').status_code,409)
    def test_last_active_superadmin_cannot_be_deleted(self):
        self.login('super@fixture.test')
        self.post('/superadmin/admins/delete/1')
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM admins WHERE id=1')[0]['total'],1)
    def test_last_windu_approver_with_pending_requests_is_protected(self):
        self.query("UPDATE pegawai SET perusahaan_induk='Windu Karya' WHERE id=1")
        self.query("UPDATE admins SET company='Windu Karya' WHERE id=2")
        self.login('super@fixture.test')
        self.post('/superadmin/admins/delete/2')
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM admins WHERE id=2')[0]['total'],1)
    def test_future_legacy_transactions_still_reserve_plafond(self):
        self.query('UPDATE transactions SET tanggal=? WHERE user_id=1',('2099-12-31',))
        with self.app.test_request_context('/dashboard'):
            from datetime import date
            limits=web.compute_limits(6000000,date.fromisoformat(TODAY),1)
        self.assertEqual(limits['total_sukses'],20000)
    def test_legacy_migration_is_disabled_by_default(self):
        import subprocess
        environment={**os.environ,'ALLOW_LEGACY_DB_REPLACE':'0'}
        response=subprocess.run([sys.executable,'-m','app.migrasi_local'],cwd=ROOT,env=environment,capture_output=True,text=True)
        self.assertNotEqual(response.returncode,0)
        self.assertIn('Dinonaktifkan secara default',response.stderr)
    def test_deployment_entrypoint_exists(self):
        import json
        config=json.loads((ROOT/'vercel.json').read_text())
        for build in config['builds']:
            self.assertTrue((ROOT/build['src']).is_file())
        self.assertEqual((ROOT/'requirements_vercel.txt').read_text().strip(),'-r requirements.txt')
    def test_queue_backend_reads_environment_after_import(self):
        from app.tasks import queue
        with patch.dict(os.environ,{'REDIS_URL':'redis://fixture.invalid/0'}):
            self.assertEqual(queue.get_mode(),'rq')
        with patch.dict(os.environ,{'REDIS_URL':''}):
            self.assertEqual(queue.get_mode(),'inline')
    def test_queue_restarts_safely_after_fork(self):
        from app.tasks import queue
        queue._reset_after_fork()
        with patch.dict(os.environ,{'REDIS_URL':''}),patch.object(queue.threading,'Thread') as thread:
            queue.start_worker(self.app)
            queue.start_worker(self.app)
            self.assertEqual(thread.call_count,1)
            queue._reset_after_fork()
            queue.start_worker(self.app)
            self.assertEqual(thread.call_count,2)
        queue._reset_after_fork()
    def test_all_templates_compile_and_static_endpoints_exist(self):
        from jinja2 import nodes
        for path in (ROOT/'templates').glob('*.html'):
            text=path.read_text(encoding='utf-8')
            template=self.app.jinja_env.parse(text)
            for call in template.find_all(nodes.Call):
                if isinstance(call.node,nodes.Name) and call.node.name=='url_for' and call.args and isinstance(call.args[0],nodes.Const):
                    self.assertIn(call.args[0].value,self.app.view_functions,path.name)
    def test_security_headers_and_secret_guard(self):
        response=self.client.get('/login')
        self.assertEqual(response.headers['X-Content-Type-Options'],'nosniff')
        self.assertIn('no-store',response.headers['Cache-Control'])
        with self.assertRaises(RuntimeError):
            create_app({**CONFIG,'DB_PATH':str(self.path),'APP_SECRET':'dev-secret-change-me'})

    def test_ppn_uses_exact_integer_rounding(self):
        with patch.object(web, 'get_ppn_enabled', return_value=True):
            for amount in (*range(10001), 15000, 50000, 100000, 30_000_000):
                self.assertEqual(web.apply_ppn(amount), amount + (amount * 11 + 99) // 100, amount)
        with patch.object(web, 'get_ppn_enabled', return_value=False):
            self.assertEqual(web.apply_ppn(10000), 10000)

    def test_withdrawal_persists_exact_fee(self):
        self.query('UPDATE pegawai SET admin_fee_flat=10000 WHERE id=1')
        self.login('employee1@fixture.test')
        data = self.withdrawal_data()
        with patch.object(web, 'enqueue_email'):
            self.assertEqual(self.post('/tarik-gaji', data).status_code, 302)
        self.assertEqual(self.query('SELECT admin_fee FROM transactions WHERE request_key IS NOT NULL')[0]['admin_fee'], 11100)

    def test_ppn_lookup_failure_cannot_silently_remove_tax(self):
        self.login('employee1@fixture.test')
        data = self.withdrawal_data()
        original = web.get_ppn_enabled
        calls = []

        def lookup():
            calls.append(True)
            if len(calls) == 1:
                return original()
            raise sqlite3.OperationalError('Simulated settings lookup failure')

        with patch.object(web, 'get_ppn_enabled', side_effect=lookup):
            with self.assertRaises(sqlite3.OperationalError):
                self.post('/tarik-gaji', data)
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM transactions WHERE request_key IS NOT NULL')[0]['total'], 0)

    def test_export_period_follows_jakarta_at_month_boundary(self):
        from datetime import date
        from app.tasks import export_scheduler

        class HostDate(date):
            @classmethod
            def today(cls):
                return cls(2026, 10, 31)

        for today, expected in ((date(2026, 11, 1), (date(2026, 11, 1), date(2026, 12, 1))),
                                (date(2027, 1, 1), (date(2027, 1, 1), date(2027, 2, 1)))):
            with patch.object(export_scheduler, 'date', HostDate), patch.object(export_scheduler, 'business_today', return_value=today):
                self.assertEqual(export_scheduler._current_period_range(), expected)

    def test_token_pages_never_cache_or_forward_token_in_referrer(self):
        for path in ('/verify-registration/invalid', '/reset-password/invalid'):
            response = self.client.get(path)
            self.assertIn('no-store', response.headers['Cache-Control'])
            self.assertEqual(response.headers['Referrer-Policy'], 'no-referrer')

    def test_explicit_test_database_overrides_environment(self):
        unused = Path(self.directory.name) / 'must-not-open-from-environment.db'
        with patch.dict(os.environ, {'DB_PATH': str(unused)}):
            app = create_app({**CONFIG, 'DB_PATH': str(self.path)})
            self.assertEqual(app.config['DB_PATH'], str(self.path))
        self.assertFalse(unused.exists())

    def test_stale_password_changes_cannot_overwrite_new_credentials(self):
        cases = (('employee1@fixture.test', '/settings', 'user_accounts', 1),
                 ('admin@example.com', '/admin/settings', 'admins', 2),
                 ('super@fixture.test', '/superadmin/settings', 'admins', 1))
        replacement = generate_password_hash('ChangedByAnotherRequest123')
        for email, path, table, account_id in cases:
            with self.subTest(path=path):
                self.client = self.app.test_client()
                self.login(email)

                def simultaneous_change(password):
                    self.query(f'UPDATE {table} SET password_hash=? WHERE id=?', (replacement, account_id))
                    return generate_password_hash(password)

                with patch.object(web, 'generate_password_hash', side_effect=simultaneous_change):
                    response = self.post(path, {'old_password': PASSWORD, 'new_password': 'StaleRequest123', 'new_password2': 'StaleRequest123'})
                self.assertEqual(response.location, '/login')
                self.assertEqual(self.query(f'SELECT password_hash FROM {table} WHERE id=?', (account_id,))[0]['password_hash'], replacement)
                with self.client.session_transaction() as session:
                    self.assertNotIn('auth_version', session)

    def test_superseded_verification_cannot_create_stale_account(self):
        import hashlib
        with patch.object(email_service, 'send_email', return_value=True) as send:
            self.client.post('/register', data={'csrf_token': self.csrf('/register'), 'name': 'New Fixture',
                'email': 'employee3@fixture.test', 'password': PASSWORD, 'password2': PASSWORD})
        path = re.search(r'https://fixture.example(/verify-registration/[^\s]+)', send.call_args.args[1]).group(1)
        self.csrf(path)

        def supersede(*args, **kwargs):
            self.query('UPDATE pending_registrations SET token_hash=? WHERE email=?',
                       (hashlib.sha256(b'replaced-token').hexdigest(), 'employee3@fixture.test'))
            return None

        with patch.object(web, 'rate_limit', side_effect=supersede):
            response = self.post(path)
        self.assertEqual(response.location, '/register')
        self.assertEqual(self.query('SELECT COUNT(*) AS total FROM user_accounts WHERE pegawai_id=3')[0]['total'], 0)

    def test_excel_entity_payload_is_rejected_without_importing(self):
        import zipfile
        import openpyxl.xml
        self.assertTrue(openpyxl.xml.DEFUSEDXML)
        self.login()
        workbook = Workbook()
        workbook.active.append(['ID', 'Nama', 'Email', 'Jabatan', 'Project', 'Rek', 'Rek2', 'EW', 'Telp', 'Gaji', 'Status', 'Siklus', 'Fee', 'Perusahaan'])
        workbook.active.append(['XML1', 'XML Fixture', 'xml@fixture.test', 'Staff', 'Project 1', '1234', '', '', '0800', 4000000, 'Aktif', 'A', 15000, 'Fixture PT One'])
        original = io.BytesIO()
        workbook.save(original)
        original.seek(0)
        malicious = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(malicious, 'w', zipfile.ZIP_DEFLATED) as target:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == 'xl/workbook.xml':
                    content = b'<!DOCTYPE workbook [<!ENTITY fixture "unsafe">]>\n' + content.replace(b'<workbookPr/>', b'<workbookPr codeName="&fixture;"/>')
                target.writestr(entry, content)
        malicious.seek(0)
        response = self.post('/admin/pegawai/import', {'file': (malicious, 'entities.xlsx')})
        self.assertEqual(response.status_code, 302)
        with self.client.session_transaction() as session:
            self.assertTrue(any('File tidak bisa dibaca' in message for category, message in session['_flashes']))
        self.assertEqual(self.query("SELECT COUNT(*) AS total FROM pegawai WHERE id_pegawai='XML1'")[0]['total'], 0)

if __name__ == '__main__':
    unittest.main(verbosity=2)
