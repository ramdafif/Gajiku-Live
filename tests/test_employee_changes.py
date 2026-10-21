import io
import json
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from test_company_lifecycle import harness, at
from app.db import get_db
from app.routes import web
from app.services import company_lifecycle as lifecycle
from app.services.employee_finance import (compute_employee_limits, financial_values, migrate_finance,
                                           monthly_start, next_start, period_bounds)


def seed_projects(harness):
    for client_id in (1, 2):
        harness.query('INSERT INTO client_projects (client_id,name,name_key,created_at,updated_at) VALUES (?,?,?,?,?)',
                      (client_id, f'Project {client_id}', f'project {client_id}', '2026-01-01', '2026-01-01'))


def employee_form(**overrides):
    return {'id_pegawai': 'E1', 'nama': 'Employee 1', 'email': 'employee1@fixture.test', 'jabatan': '',
            'perusahaan_induk': 'Fixture PT One', 'perusahaan': 'Project 1', 'no_rekening': '123456',
            'status': '1', 'gaji': '6000000', 'plafon': '3000000', 'plafon_mode': 'nominal',
            'siklus_gaji': 'A', **overrides}


def pending(harness):
    return harness.query("SELECT * FROM employee_change_items WHERE status='pending' ORDER BY id")


def decide(harness, item, decision='approve', **extra):
    return harness.post(f"/employee-changes/{item['batch_id']}",
                        {'item_ids': [str(item['id'])], 'decision': decision, **extra})


def test_financial_change_waits_but_personal_fields_apply(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(nama='Udin', gaji='8000000',
                 plafon_mode='persentase', plafon_persen='50'))
    row = harness.query('SELECT nama,gaji,plafon FROM pegawai WHERE id=1')[0]
    assert row == {'nama': 'Udin', 'gaji': 6000000, 'plafon': 3000000}
    item = pending(harness)[0]
    harness.login()
    assert decide(harness, item).status_code == 302
    row = harness.query('SELECT nama,gaji,plafon FROM pegawai WHERE id=1')[0]
    assert row == {'nama': 'Udin', 'gaji': 8000000, 'plafon': 4000000}
    decide(harness, item)
    assert harness.query("SELECT COUNT(*) AS total FROM employee_audit WHERE action='approved'")[0]['total'] == 1


def test_lowering_also_waits_and_rejection_needs_reason(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(plafon='1000000'))
    item = pending(harness)[0]
    assert harness.query('SELECT plafon FROM pegawai WHERE id=1')[0]['plafon'] == 3000000
    harness.login()
    decide(harness, item, 'reject')
    assert pending(harness)
    decide(harness, item, 'reject', reason='Periksa MoU')
    assert not pending(harness)
    assert harness.query('SELECT plafon FROM pegawai WHERE id=1')[0]['plafon'] == 3000000


def test_create_is_reserved_and_invisible_until_approved(harness):
    seed_projects(harness)
    harness.login('admin@fixture.test')
    values = employee_form(id_pegawai='NEW1', email='new1@fixture.test', gaji='', plafon='1000000')
    harness.post('/admin/pegawai/add', values)
    assert not harness.query("SELECT id FROM pegawai WHERE id_pegawai='NEW1'")
    item = pending(harness)[0]
    harness.post('/admin/pegawai/add', values)
    assert len(pending(harness)) == 1
    assert len(harness.query('SELECT * FROM employee_identity_reservations')) == 2
    harness.login()
    decide(harness, item)
    assert harness.query("SELECT gaji,plafon,admin_fee_flat FROM pegawai WHERE id_pegawai='NEW1'")[0] == {
        'gaji': None, 'plafon': 1000000, 'admin_fee_flat': None}
    assert not harness.query('SELECT * FROM employee_identity_reservations')


def test_cancellation_releases_identity(harness):
    seed_projects(harness)
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/add', employee_form(id_pegawai='NEW1', email='new1@fixture.test'))
    item = pending(harness)[0]
    decide(harness, item, 'cancel')
    assert not pending(harness)
    assert not harness.query('SELECT * FROM employee_identity_reservations')


def test_tenant_access_and_fee_payload_are_denied(harness):
    harness.login('admin@fixture.test')
    response = harness.post('/admin/pegawai/1/update', employee_form(admin_fee_flat='1'))
    assert response.status_code == 403
    harness.post('/admin/pegawai/2/update', employee_form(gaji='9999999'))
    assert harness.query('SELECT gaji FROM pegawai WHERE id=2')[0]['gaji'] == 6000000
    harness.post('/admin/pegawai/1/update', employee_form(gaji='7000000'))
    item = pending(harness)[0]
    assert decide(harness, item).status_code == 403
    harness.login('other@fixture.test')
    assert harness.client.get(f"/employee-changes/{item['batch_id']}").status_code in (403, 404)


def test_deactivate_cancels_pending_and_activation_waits(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(gaji='7000000'))
    harness.post('/admin/pegawai/1/update', employee_form(status='0'))
    assert not pending(harness)
    assert harness.query('SELECT status_aktif FROM pegawai WHERE id=1')[0]['status_aktif'] == 0
    harness.post('/admin/pegawai/1/update', employee_form(status='1'))
    assert harness.query('SELECT status_aktif FROM pegawai WHERE id=1')[0]['status_aktif'] == 0
    assert pending(harness)


def test_conflicting_approval_does_not_overwrite_new_values(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(gaji='7000000'))
    item = pending(harness)[0]
    harness.query('UPDATE pegawai SET gaji=9000000 WHERE id=1')
    harness.login()
    decide(harness, item)
    assert harness.query('SELECT status FROM employee_change_items WHERE id=?', (item['id'],))[0]['status'] == 'conflict'
    assert harness.query('SELECT gaji FROM pegawai WHERE id=1')[0]['gaji'] == 9000000


def test_email_and_bank_change_preserve_account_and_transactions(harness):
    harness.query("INSERT INTO transactions (user_id,tanggal,periode,nominal,admin_fee,status,created_at,rekening_tujuan) VALUES (1,'2026-10-19','2026-10',700000,15000,'sukses','2026-10-19','OLD')")
    with harness.app.app_context():
        before = compute_employee_limits(date(2026, 10, 19), 1)['total_sukses']
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(email='changed@fixture.test', no_rekening='999999'))
    assert not pending(harness)
    assert harness.query('SELECT pegawai_id,email FROM users WHERE id=1')[0] == {'pegawai_id': 1, 'email': 'changed@fixture.test'}
    assert harness.query('SELECT rekening_tujuan FROM transactions WHERE user_id=1')[0]['rekening_tujuan'] == 'OLD'
    with harness.app.app_context():
        assert compute_employee_limits(date(2026, 10, 19), 1)['total_sukses'] == before
    assert any('Tujuan pencairan berubah' in row['title'] for row in harness.query('SELECT title FROM notifications'))


def test_fee_default_and_override(harness):
    harness.query('UPDATE pegawai SET admin_fee_flat=NULL WHERE id=1')
    harness.query('UPDATE clients SET default_admin_fee=17000 WHERE id=1')
    with harness.app.app_context():
        assert web.get_admin_fee_flat_for_user(1) == 17000
    harness.query('UPDATE pegawai SET admin_fee_flat=0 WHERE id=1')
    with harness.app.app_context():
        assert web.get_admin_fee_flat_for_user(1) == 0


def test_import_template_is_scoped_and_creates_pending_batch(harness):
    seed_projects(harness)
    harness.login('admin@fixture.test')
    response = harness.client.get('/admin/pegawai/import/template')
    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert 'Admin Fee' not in headers and 'Perusahaan*' not in headers and 'Project*' in headers
    data = {'ID Pegawai*': 'IMP1', 'Nama*': 'Imported', 'Email*': 'imported@fixture.test',
            'Mode Plafon*': 'nominal', 'Plafon Nominal': 1000000, 'Project*': 'Project 1'}
    sheet.append([data.get(header) for header in headers])
    data.update({'ID Pegawai*': 'IMP2', 'Email*': 'imported2@fixture.test'})
    sheet.append([data.get(header) for header in headers])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    harness.post('/admin/pegawai/import', {'file': (output, 'import.xlsx')})
    assert not harness.query("SELECT id FROM pegawai WHERE id_pegawai LIKE 'IMP%'")
    items = pending(harness)
    assert len(items) == 2 and items[0]['batch_id'] == items[1]['batch_id']
    harness.login()
    decide(harness, items[0])
    assert len(pending(harness)) == 1
    assert harness.query("SELECT id FROM pegawai WHERE id_pegawai='IMP1'")


def test_nonproject_template_hides_project(harness):
    harness.query('UPDATE clients SET allow_projects=0 WHERE id=1')
    harness.login('admin@fixture.test')
    workbook = load_workbook(io.BytesIO(harness.client.get('/admin/pegawai/import/template').data))
    assert not any('Project' in str(cell.value) for cell in workbook.active[1])
    workbook.close()

def test_import_reports_every_invalid_row_without_filling_session(harness):
    harness.login('admin@fixture.test')
    workbook = load_workbook(io.BytesIO(harness.client.get('/admin/pegawai/import/template').data))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    for index in range(25):
        row = {'ID Pegawai*': f'INVALID{index}', 'Nama*': 'Invalid', 'Email*': 'not-an-email'}
        sheet.append([row.get(header) for header in headers])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    response = harness.post('/admin/pegawai/import', {'file': (output, 'invalid.xlsx')})
    assert response.status_code == 200
    page = response.get_data(as_text=True)
    assert all(f'Baris {number}:' in page for number in range(2, 27))
    assert not pending(harness)
    with harness.client.session_transaction() as state:
        assert not any('Baris ' in message for category, message in state.get('_flashes', []))


def test_bulk_preview_freezes_targets_and_queues_changes(harness):
    harness.login('admin@fixture.test')
    response = harness.post('/admin/pegawai/bulk', {'client': 'Fixture PT One', 'all_filtered': '1',
                            'set_plafon': '1', 'plafon_mode': 'nominal', 'plafon': '1000000'})
    assert response.status_code == 200
    token = re.search(r'name="preview_token" value="([^"]+)"', response.get_data(as_text=True)).group(1)
    harness.post('/admin/pegawai/bulk', {'preview_token': token})
    assert len(pending(harness)) == 3
    assert harness.query('SELECT plafon FROM pegawai WHERE id=1')[0]['plafon'] == 3000000
    harness.post('/admin/pegawai/bulk', {'preview_token': token})
    assert len(pending(harness)) == 3


def test_bulk_preview_conflict_rolls_back(harness):
    harness.login()
    response = harness.post('/admin/pegawai/bulk', {'client': 'Fixture PT One', 'employee_ids': ['1'],
                            'set_plafon': '1', 'plafon_mode': 'nominal', 'plafon': '1000000'})
    token = re.search(r'name="preview_token" value="([^"]+)"', response.get_data(as_text=True)).group(1)
    harness.query("UPDATE pegawai SET nama='Changed' WHERE id=1")
    harness.post('/admin/pegawai/bulk', {'preview_token': token})
    assert harness.query('SELECT plafon FROM pegawai WHERE id=1')[0]['plafon'] == 3000000


@pytest.mark.parametrize('path', ['/reset', '/admin/reset', '/admin/reset_all', '/superadmin/reset', '/superadmin/reset_all'])
def test_destructive_reset_removed_even_when_flag_enabled(harness, path):
    harness.login()
    harness.app.config['ALLOW_DESTRUCTIVE_RESET'] = True
    assert harness.post(path).status_code == 404
    assert len(harness.query('SELECT id FROM pegawai')) == 4


def configure_weekly(harness):
    harness.query("UPDATE clients SET plafon_start_day=16,plafon_confirmed_at='2026-10-01',period_history=? WHERE id=1",
                  (json.dumps([{'from': '2026-09-16', 'day': 16}]),))
    harness.query('UPDATE admins SET cutoff_mingguan_aktif=1,cutoff_hari=2 WHERE company=?', ('Fixture PT One',))
    harness.query('UPDATE pegawai SET gaji=NULL,plafon=4000000 WHERE id=1')


def transaction(harness, day, amount, status='sukses'):
    harness.query('''INSERT INTO transactions (user_id,tanggal,periode,nominal,admin_fee,status,created_at)
        VALUES (1,?,'legacy-key',?,15000,?,?)''', (day, amount, status, day))


def test_week_crosses_month_without_new_weekly_allowance(harness):
    configure_weekly(harness)
    transaction(harness, '2026-10-15', 700000)
    with harness.app.app_context():
        limits = compute_employee_limits(date(2026, 10, 16), 1)
    assert limits['sisa_plafon'] == 300000
    assert limits['total_bulanan'] == 0
    assert limits['hari_ke'] == 2
    assert limits['gaji_pokok'] is None


def test_fifth_week_and_pending_requests_cannot_exceed_month(harness):
    configure_weekly(harness)
    for day in ('2026-10-16', '2026-10-23', '2026-10-30'):
        transaction(harness, day, 1000000)
    transaction(harness, '2026-11-06', 1000000, 'on-proses')
    with harness.app.app_context():
        limits = compute_employee_limits(date(2026, 11, 13), 1)
    assert limits['sisa_bulanan'] == 0
    assert limits['sisa_plafon'] == 0 and limits['saldo'] == 0


def test_percentage_requires_salary_and_recomputes(harness):
    assert financial_values({'gaji': '6000001', 'plafon_mode': 'persentase', 'plafon_persen': '50'})['plafon'] == 3000000
    with pytest.raises(ValueError):
        financial_values({'gaji': '', 'plafon_mode': 'persentase', 'plafon_persen': '50'})


def test_migration_dry_run_and_idempotence(harness):
    harness.query('UPDATE pegawai SET plafon=NULL')
    with harness.app.app_context():
        db = get_db()
        assert migrate_finance(db)['employees'] == 4
        assert db.execute('SELECT gaji FROM pegawai WHERE id=1').fetchone()['gaji'] == 6000000
        migrate_finance(db, True)
        assert db.execute('SELECT gaji FROM pegawai WHERE id=1').fetchone()['gaji'] is None
        assert db.execute('SELECT plafon FROM pegawai WHERE id=1').fetchone()['plafon'] == 3000000
        assert migrate_finance(db, True)['already_migrated']


def test_unlinked_legacy_user_blocks_migration(harness):
    harness.query("UPDATE users SET email='orphan@fixture.test',pegawai_id=NULL WHERE id=1")
    with harness.app.app_context():
        with pytest.raises(ValueError):
            migrate_finance(get_db(), True)


@pytest.mark.parametrize('year,month,last', [(2026, 2, 28), (2028, 2, 29), (2026, 4, 30)])
def test_period_day_31_clamps_without_drift(year, month, last):
    start = monthly_start(date(year, month, last), 31)
    assert start == date(year, month, last)
    assert next_start(start, 31).day == 31


def test_monthly_billing_blocks_at_boundary_not_due_date(harness, monkeypatch):
    configure_weekly(harness)
    harness.query('UPDATE clients SET period_history=? WHERE id=1',
                  (json.dumps([{'from': '2026-10-16', 'day': 16}]),))
    with harness.app.app_context():
        db = get_db()
        with lifecycle.locked_client(db, 1) as client:
            lifecycle.configure_billing(db, client, 20, 'fixture', at(10, 19))
    assert not harness.state(at(10, 19))['blocked']
    assert not harness.state(at(10, 26))['blocked']
    state = harness.state(at(11, 16))
    assert state['blocked']
    bill = next(bill for bill in state['overdue'] if bill['service_start'] == '2026-10-16')
    assert str(bill['due_on']) == '2026-11-20'
    harness.action(lifecycle.mark_paid, bill['id'], now=at(11, 16))
    assert harness.state(at(11, 16))['blocked']
    harness.action(lifecycle.activate_company, now=at(11, 16))
    assert not harness.state(at(11, 16))['blocked']
    assert harness.query('SELECT status_aktif FROM pegawai WHERE id=3')[0]['status_aktif'] == 0


def test_personal_change_after_request_is_not_overwritten(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(gaji='7000000'))
    item = pending(harness)[0]
    harness.post('/admin/pegawai/1/update', employee_form(nama='Nama terbaru'))
    harness.login()
    decide(harness, item)
    assert harness.query('SELECT nama,gaji FROM pegawai WHERE id=1')[0] == {'nama': 'Nama terbaru', 'gaji': 7000000}


def test_two_superadmins_cannot_approve_twice(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(gaji='7000000'))
    item = pending(harness)[0]
    clients = [harness.login(client=harness.app.test_client()) for index in range(2)]
    def approve(client):
        return harness.post(f"/employee-changes/{item['batch_id']}",
                            {'decision': 'approve', 'item_ids': [str(item['id'])]}, client).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert list(pool.map(approve, clients)) == [302, 302]
    assert harness.query("SELECT COUNT(*) AS total FROM employee_audit WHERE action='approved'")[0]['total'] == 1


def test_all_filtered_selection_includes_other_pages(harness):
    with harness.app.app_context():
        db = get_db()
        for index in range(55):
            db.execute('''INSERT INTO pegawai (id_pegawai,nama,email,gaji,plafon,status_aktif,
                perusahaan_induk,perusahaan,siklus_gaji,created_at) VALUES (?,?,?,6000000,3000000,1,
                'Fixture PT One','Project 1','A','2026-10-19')''',
                (f'B{index}', f'Bulk {index}', f'bulk{index}@fixture.test'))
        db.commit()
    harness.login('admin@fixture.test')
    assert harness.client.get('/admin/pegawai').get_data(as_text=True).count('name="employee_ids"') == 50
    response = harness.post('/admin/pegawai/bulk', {'client': 'Fixture PT One', 'all_filtered': '1',
                            'set_plafon': '1', 'plafon_mode': 'nominal', 'plafon': '1000000'})
    token = re.search(r'name="preview_token" value="([^"]+)"', response.get_data(as_text=True)).group(1)
    harness.post('/admin/pegawai/bulk', {'preview_token': token})
    assert len(pending(harness)) == 58


def test_revoked_submitter_requires_new_request(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(gaji='7000000'))
    item = pending(harness)[0]
    harness.login()
    harness.query('UPDATE admins SET status_aktif=0 WHERE id=2')
    decide(harness, item)
    assert pending(harness)
    assert harness.query('SELECT gaji FROM pegawai WHERE id=1')[0]['gaji'] == 6000000


def test_policy_and_cutoff_changes_take_effect_after_current_period(harness):
    configure_weekly(harness)
    harness.login()
    harness.post('/superadmin/clients/1/employee-policy', {'plafon_start_day': '1', 'default_admin_fee': '17000'})
    with harness.app.app_context():
        from app.services.employee_finance import schedule_cutoff, weekly_config
        db = get_db()
        client = dict(db.execute('SELECT * FROM clients WHERE id=1').fetchone())
        assert period_bounds(client, date(2026, 10, 20)) == (date(2026, 10, 16), date(2026, 11, 16))
        assert period_bounds(client, date(2026, 11, 17)) == (date(2026, 11, 16), date(2026, 12, 1))
        schedule_cutoff(db, client, 1, 4)
        db.commit()
        assert weekly_config(db, client['name'], date(2026, 10, 21))['cutoff_hari'] == 2
        assert weekly_config(db, client['name'], date(2026, 10, 22))['cutoff_hari'] == 4


def test_legacy_bills_require_explicit_mapping(harness):
    configure_weekly(harness)
    harness.query('UPDATE clients SET period_history=? WHERE id=1',
                  (json.dumps([{'from': '2026-10-16', 'day': 16}]),))
    harness.query("INSERT INTO client_billing (client_id,due_day,first_due,updated_by,updated_at) VALUES (1,20,'2026-11-20','fixture','2026-10-19')")
    harness.query("INSERT INTO client_bills (client_id,period,due_on,created_at) VALUES (1,'2026-10','2026-11-20','2026-10-19')")
    state = harness.state(at(10, 19))
    assert state['configuration_required'] and state['blocked']
    bill_id = state['legacy_bill_ids'][0]
    runner = harness.app.test_cli_runner()
    args = ['map-legacy-bill', '--bill-id', str(bill_id), '--service-start', '2026-10-16', '--service-end', '2026-11-16']
    assert runner.invoke(args=args).exit_code == 0
    assert harness.state(at(10, 19))['blocked']
    assert runner.invoke(args=[*args, '--apply']).exit_code == 0
    assert not harness.state(at(10, 19))['blocked']


def test_management_and_review_pages_render(harness):
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(gaji='7000000'))
    item = pending(harness)[0]
    for path in ('/admin/pegawai', '/employee-changes', f"/employee-changes/{item['batch_id']}"):
        assert harness.client.get(path).status_code == 200
    harness.login()
    assert harness.client.get('/superadmin/clients/1').status_code == 200
    assert harness.client.get(f"/employee-changes/{item['batch_id']}").status_code == 200

def test_weekly_kpi_uses_effective_history_without_counting_admins(harness):
    configure_weekly(harness)
    harness.query('UPDATE clients SET cutoff_history=? WHERE id=1', (json.dumps([
        {'from': '1900-01-01', 'enabled': 1, 'day': 2},
        {'from': '2026-10-22', 'enabled': 1, 'day': 4},
    ]),))
    with harness.app.app_context():
        assert web.get_weekly_cutoff_counts(get_db()) == {2: 2}

def test_history_uses_employee_identity_after_email_change(harness):
    harness.query("INSERT INTO transactions (user_id,tanggal,periode,nominal,admin_fee,status,created_at,rekening_tujuan) VALUES (1,'2026-10-19','2026-10',700000,15000,'sukses','2026-10-19','ORIGINAL-DESTINATION')")
    harness.login('admin@fixture.test')
    harness.post('/admin/pegawai/1/update', employee_form(email='changed@fixture.test', no_rekening='NEW-DESTINATION'))
    for path in ('/admin/riwayat', '/superadmin/riwayat'):
        if path.startswith('/superadmin'):
            harness.login()
        response = harness.client.get(path)
        assert response.status_code == 200
        assert 'ORIGINAL-DESTINATION' in response.get_data(as_text=True)
