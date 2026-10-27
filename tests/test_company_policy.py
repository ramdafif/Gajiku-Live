import io
import json
import re
import shutil
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from flask import template_rendered
from openpyxl import Workbook

import test_audit_regression as regression
from app import create_app
from app.db import get_db
from app.routes import clients, web
from app.services.company_policy import (ensure_company_policy_schema, effective_reg_limit,
                                         guard_cutoff_change, lock_policy, policy_period)
from app.utils.cache import _cache


@pytest.fixture(scope="module")
def policy_base():
    regression.AuditRegression.setUpClass()
    with tempfile.TemporaryDirectory(prefix="gajiku-policy-base-") as directory:
        base = Path(directory) / "base.db"
        shutil.copyfile(regression.AuditRegression.base, base)
        regression.AuditRegression.tearDownClass()
        yield base


@pytest.fixture
def harness(policy_base, tmp_path):
    _cache.clear()
    instance = regression.AuditRegression()
    instance.path = tmp_path / "policy.db"
    shutil.copyfile(policy_base, instance.path)
    instance.app = create_app({**regression.CONFIG, "DB_PATH": str(instance.path)})
    instance.client = instance.app.test_client()
    yield instance
    instance.doCleanups()


def save_policy(harness, section, client=None, **values):
    version = harness.query("SELECT policy_version FROM clients WHERE id=1")[0]["policy_version"]
    return harness.post("/superadmin/clients/1", {"section": section, "policy_version": str(version), **values}, client)


def employee_data(harness, employee_id=1, **values):
    row = harness.query("SELECT * FROM pegawai WHERE id=?", (employee_id,))[0]
    data = {field: str(row[field] or "") for field in ("id_pegawai", "nama", "email", "jabatan", "gaji", "perusahaan_induk", "perusahaan", "no_rekening", "siklus_gaji", "admin_fee_flat")}
    data["status"] = str(row["status_aktif"])
    return {**data, **values}


def test_weekly_migration_matches_legacy_on_same_dates_and_keeps_history(harness):
    harness.query("DELETE FROM app_settings WHERE `key`='company_policy_v1'")
    harness.query("UPDATE admins SET cutoff_mingguan_aktif=1,cutoff_hari=4,produk_reg_aktif=0 WHERE id=2")
    harness.query("UPDATE admins SET cutoff_mingguan_aktif=0,cutoff_hari=6,produk_reg_aktif=1,produk_urg_aktif=0 WHERE id=4")
    harness.query("UPDATE pegawai SET siklus_gaji='D' WHERE id=1")
    days = [date(2026, 12, 24) + timedelta(days=offset) for offset in range(24)]
    days += [date(2028, 2, 25) + timedelta(days=offset) for offset in range(10)]
    for offset, day in enumerate(days):
        harness.query("INSERT INTO transactions(user_id,tanggal,periode,nominal,admin_fee,status,created_at,product) VALUES (1,?,?,1234,50,?,?,?)",
                      (day.isoformat(), web.weekly_period_key(day, 4), "sukses" if offset % 2 else "on-proses", day.isoformat(), "urg" if offset % 2 else "reg"))
    history = harness.query("SELECT * FROM transactions ORDER BY id")

    def legacy(company):
        row = get_db().execute("SELECT cutoff_hari FROM admins WHERE LOWER(TRIM(company))=LOWER(TRIM(?)) AND COALESCE(cutoff_mingguan_aktif,0)=1 ORDER BY id DESC LIMIT 1", (company,)).fetchone()
        if not row:
            return None
        day = int(row["cutoff_hari"] if row["cutoff_hari"] is not None else 2)
        return {"cutoff_hari": day if day in range(7) else 2}

    with harness.app.app_context():
        with patch.object(web, "get_weekly_cutoff_config", legacy):
            before = [web.compute_limits(6000000, day, employee_id) for day in days for employee_id in (1, 2)]
        ensure_company_policy_schema(get_db(), "sqlite")
        after = [web.compute_limits(6000000, day, employee_id) for day in days for employee_id in (1, 2)]
        assert before == after
        assert web.get_weekly_cutoff_config(" FIXTURE PT ONE ") == {"cutoff_hari": 4}
        assert web.get_products_for_company("Fixture PT One") == ["reg"]
        assert effective_reg_limit(get_db(), "Fixture PT One")["effective"] == 1000000
    assert harness.query("SELECT * FROM transactions ORDER BY id") == history
    assert all(row["reg_withdrawal_max_override"] is None for row in harness.query("SELECT reg_withdrawal_max_override FROM pegawai"))
    event = harness.query("SELECT before_json FROM company_policy_events WHERE kind='migration' AND client_id=1")[-1]
    assert len(json.loads(event["before_json"])["legacy"]) == 2
    harness.query("UPDATE clients SET cutoff_hari=1 WHERE id=1")
    with harness.app.app_context():
        ensure_company_policy_schema(get_db(), "sqlite")
        assert web.get_weekly_cutoff_config("Fixture PT One") == {"cutoff_hari": 1}


@pytest.mark.parametrize("cutoff", range(7))
def test_period_boundaries_match_existing_helpers(cutoff):
    for anchor in (date(2026, 12, 28), date(2028, 2, 27)):
        for offset in range(10):
            day = anchor + timedelta(days=offset)
            start, end = policy_period(day, True, cutoff, "D")
            assert start == web.weekly_period_start(day, cutoff)
            assert start <= day < end
            assert (end - start).days == 7
            assert web.weekly_day_in_period(day, cutoff) == (day - start).days + 1
            for cycle in "ABCD":
                monthly_start, monthly_end = policy_period(day, False, cutoff, cycle)
                assert web.period_key_by_cycle(day, cycle) == monthly_start.strftime("%Y-%m")
                assert monthly_start <= day < monthly_end


def test_cap_only_never_changes_weekly_values_or_history_and_profile_is_separate(harness):
    harness.query("UPDATE clients SET cutoff_mingguan_aktif=1,cutoff_hari=2 WHERE id=1")
    history = harness.query("SELECT * FROM transactions ORDER BY id")
    with harness.app.app_context():
        before = web.compute_limits(6000000, date(2026, 10, 27), 1)
    harness.login("super@fixture.test")
    assert save_policy(harness, "reg_limit", reg_withdrawal_max="500000", cutoff_hari="6").status_code == 302
    with harness.app.app_context():
        assert web.compute_limits(6000000, date(2026, 10, 27), 1) == before
    assert harness.post("/superadmin/clients/1", {"name": "Fixture PT One", "allow_projects": "1", "policy_version": "2", "reg_withdrawal_max": "9999"}).status_code == 302
    row = harness.query("SELECT * FROM clients WHERE id=1")[0]
    assert row["reg_withdrawal_max"] == 500000
    assert row["cutoff_hari"] == 2
    assert harness.query("SELECT * FROM transactions ORDER BY id") == history


def test_stale_and_missing_policy_versions_are_rejected(harness):
    harness.login("super@fixture.test")
    assert save_policy(harness, "reg_limit", reg_withdrawal_max="750000").status_code == 302
    for version in (None, "1"):
        data = {"section": "cutoff", "cutoff_mingguan_aktif": "1", "cutoff_hari": "4"}
        if version:
            data["policy_version"] = version
        assert harness.post("/superadmin/clients/1", data).status_code == 409
    assert harness.query("SELECT reg_withdrawal_max,cutoff_mingguan_aktif FROM clients WHERE id=1")[0] == {"reg_withdrawal_max": 750000, "cutoff_mingguan_aktif": 0}


def test_cutoff_guard_pending_success_old_new_windows_and_unrelated_tenant(harness):
    harness.login("super@fixture.test")
    assert save_policy(harness, "cutoff", cutoff_mingguan_aktif="1", cutoff_hari="2").status_code == 422
    harness.query("DELETE FROM transactions WHERE user_id=1")
    day = date.fromisoformat(regression.TODAY)
    harness.query("INSERT INTO transactions(user_id,tanggal,periode,nominal,admin_fee,status,created_at) VALUES (1,?,?,5000,0,'sukses',?)", (day.isoformat(), day.strftime("%Y-%m"), day.isoformat()))
    assert save_policy(harness, "cutoff", cutoff_mingguan_aktif="1", cutoff_hari="2").status_code == 422
    harness.query("UPDATE transactions SET tanggal='2020-01-01',periode='2020-01' WHERE user_id=1")
    assert save_policy(harness, "cutoff", cutoff_mingguan_aktif="1", cutoff_hari="2").status_code == 302
    assert harness.query("SELECT cutoff_mingguan_aktif FROM clients WHERE id=1")[0]["cutoff_mingguan_aktif"] == 1
    harness.query("UPDATE transactions SET status='on-proses' WHERE user_id=1")
    assert save_policy(harness, "cutoff", cutoff_mingguan_aktif="1", cutoff_hari="2").status_code == 302
    assert save_policy(harness, "cutoff", cutoff_mingguan_aktif="1", cutoff_hari="3").status_code == 422


@pytest.mark.parametrize("old_weekly,new_weekly,new_day", [(1, 0, 2), (0, 1, 4)])
def test_success_in_only_old_or_only_new_period_blocks_cutoff(harness, old_weekly, new_weekly, new_day):
    harness.query("DELETE FROM transactions")
    harness.query("UPDATE clients SET cutoff_mingguan_aktif=?,cutoff_hari=2 WHERE id=1", (old_weekly,))
    harness.query("UPDATE pegawai SET siklus_gaji=' d ' WHERE id=1")
    harness.query("INSERT INTO transactions(user_id,tanggal,periode,nominal,admin_fee,status,created_at) VALUES (1,'2026-12-27','2026-12',1000,0,'sukses','2026-12-27')")
    with harness.app.app_context():
        db = get_db()
        old = db.execute("SELECT * FROM clients WHERE id=1").fetchone()
        with pytest.raises(ValueError, match="transaksi selesai"):
            guard_cutoff_change(db, "Fixture PT One", old,
                {"cutoff_mingguan_aktif": new_weekly, "cutoff_hari": new_day}, today=date(2027, 1, 2))


def test_inactive_cutoff_day_change_does_not_trigger_guard(harness):
    harness.login("super@fixture.test")
    assert save_policy(harness, "cutoff", cutoff_mingguan_aktif="0", cutoff_hari="6").status_code == 302


def test_override_inherit_equal_default_inactive_other_tenant_and_panel(harness):
    harness.login("super@fixture.test")
    assert "Semua karyawan mengikuti batas perusahaan." in harness.client.get("/superadmin/clients/1").get_data(as_text=True)
    for employee_id, amount in ((1, "500000"), (3, "1500000"), (2, "888888")):
        assert harness.post(f"/admin/pegawai/{employee_id}/update", employee_data(harness, employee_id, reg_limit_mode="override", reg_withdrawal_max_override=amount)).status_code == 302
    harness.query("UPDATE pegawai SET status_aktif=0 WHERE id=3")
    html = harness.client.get("/superadmin/clients/1").get_data(as_text=True)
    assert "! 2 karyawan menggunakan batas khusus" in html
    assert "Employee Fixture 1" in html and "Employee Fixture 3" in html
    assert "Employee Fixture 2" not in html
    assert "Tidak aktif" in html and "client=Fixture+PT+One" in html
    assert save_policy(harness, "reg_limit", reg_withdrawal_max="750000").status_code == 302
    with harness.app.app_context():
        db = get_db()
        employees = db.execute("SELECT * FROM pegawai ORDER BY id").fetchall()
        assert effective_reg_limit(db, "Fixture PT One", employees[0])["effective"] == 500000
        assert effective_reg_limit(db, "Fixture PT One", employees[2])["effective"] == 1500000
    assert harness.post("/admin/pegawai/1/update", employee_data(harness, reg_limit_mode="override", reg_withdrawal_max_override="750000")).status_code == 302
    assert "! 2 karyawan" in harness.client.get("/superadmin/clients/1").get_data(as_text=True)
    assert harness.post("/admin/pegawai/1/update", employee_data(harness, reg_limit_mode="inherit")).status_code == 302
    assert harness.query("SELECT reg_withdrawal_max_override FROM pegawai WHERE id=1")[0]["reg_withdrawal_max_override"] is None
    assert "! 1 karyawan" in harness.client.get("/superadmin/clients/1").get_data(as_text=True)


def test_old_employee_forms_preserve_override_and_salary_cycle(harness):
    harness.query("UPDATE pegawai SET reg_withdrawal_max_override=123456,siklus_gaji='D' WHERE id=1")
    harness.login("super@fixture.test")
    data = employee_data(harness)
    del data["siklus_gaji"]
    assert harness.post("/admin/pegawai/1/update", data).status_code == 302
    assert harness.query("SELECT reg_withdrawal_max_override,siklus_gaji FROM pegawai WHERE id=1")[0] == {"reg_withdrawal_max_override": 123456, "siklus_gaji": "D"}


@pytest.mark.parametrize("raw", ["0", "-1", "1.5", "1e6", "1.000.000", "2147483648", ""])
def test_invalid_override_keeps_input_in_response_not_cookie(harness, raw):
    harness.login("super@fixture.test")
    response = harness.post("/admin/pegawai/1/update", employee_data(harness, reg_limit_mode="override", reg_withdrawal_max_override=raw, no_rekening="00123456"))
    assert response.status_code == 422
    assert "data-restore-form" in response.get_data(as_text=True)
    assert "00123456" in response.get_data(as_text=True)
    assert harness.query("SELECT reg_withdrawal_max_override FROM pegawai WHERE id=1")[0]["reg_withdrawal_max_override"] is None
    with harness.client.session_transaction() as session:
        assert "00123456" not in str(dict(session))
    assert "no-store" in response.headers["Cache-Control"]


def test_admin_cannot_overpost_override_and_does_not_change_company_policy(harness):
    harness.query("UPDATE clients SET cutoff_mingguan_aktif=1,cutoff_hari=4 WHERE id=1")
    before = harness.query("SELECT * FROM clients ORDER BY id")
    harness.query("UPDATE admins SET status_aktif=1 WHERE id=4")
    harness.login("admin@example.com")
    assert harness.post("/admin/pegawai/1/update", employee_data(harness, reg_limit_mode="override", reg_withdrawal_max_override="900000")).status_code == 422
    assert harness.query("SELECT reg_withdrawal_max_override FROM pegawai WHERE id=1")[0]["reg_withdrawal_max_override"] is None
    assert harness.client.get("/superadmin/system-settings").status_code == 403
    harness.login("super@fixture.test")
    assert harness.post("/superadmin/admins/edit/2", {"nama": "Admin New", "email": "admin@example.com", "perusahaan": "Fixture PT Two", "status": "1", "hak_approval": "1", "cutoff_hari": "6"}).status_code == 302
    assert harness.query("SELECT * FROM clients ORDER BY id") == before
    html = harness.client.get("/superadmin/admins").get_data(as_text=True)
    assert 'name="cutoff_hari"' not in html and 'name="produk_reg_aktif"' not in html


def test_override_above_company_cap_balance_and_stale_form(harness):
    harness.prepare_company_limit_withdrawal(500000)
    harness.query("UPDATE pegawai SET reg_withdrawal_max_override=1500000 WHERE id=1")
    assert harness.company_withdrawal_config()["regMax"] == 1500000
    data = harness.withdrawal_data()
    harness.query("UPDATE pegawai SET reg_withdrawal_max_override=750000 WHERE id=1")
    response = harness.post("/tarik-gaji", {**data, "nominal": "1000000"})
    assert harness.company_withdrawal_config(response)["regMax"] == 750000
    assert not harness.query("SELECT id FROM transactions")
    data = harness.withdrawal_data()
    assert harness.post("/tarik-gaji", {**data, "nominal": "750000"}).status_code == 302
    assert harness.query("SELECT nominal FROM transactions")[0]["nominal"] == 750000


@pytest.mark.parametrize("value,expected", [(None, None), ("IKUTI_PT", None), (1500000, 1500000)])
def test_import_optional_override_and_old_template(harness, value, expected):
    harness.login("super@fixture.test")
    workbook = Workbook()
    workbook.active.append(["Column"] * (15 if value is not None else 14))
    row = ["EXTRA1", "Extra Employee", "extra@fixture.test", "Staff", "Project 1", "00123", "", "", "0800", 5000000, "Aktif", "D", 15000, "Fixture PT One"]
    if value is not None:
        row.append(value)
    workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    assert harness.post("/admin/pegawai/import", {"file": (output, "import.xlsx")}).status_code == 302
    saved = harness.query("SELECT reg_withdrawal_max_override FROM pegawai WHERE id_pegawai='EXTRA1'")
    assert saved == [{"reg_withdrawal_max_override": expected}]


def test_company_can_disable_both_products_without_changing_history(harness):
    harness.login("super@fixture.test")
    before = harness.query("SELECT * FROM transactions ORDER BY id")
    assert save_policy(harness, "products", produk_reg_aktif="0", produk_urg_aktif="0").status_code == 302
    with harness.app.app_context():
        assert web.get_products_for_company("Fixture PT One") == []
    assert harness.query("SELECT * FROM transactions ORDER BY id") == before


def test_admin_import_template_has_no_default_override_marker(harness):
    from openpyxl import load_workbook
    harness.login("admin@example.com")
    response = harness.client.get("/admin/pegawai/import/template")
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.active.cell(2, 15).value is None
    workbook.close()
    workbook = Workbook()
    workbook.active.append(["Column"] * 15)
    workbook.active.append(["FORGED1", "Forged Limit", "forged@fixture.test", "Staff", "Project 1", "1234", "", "", "0800", 5000000, "Aktif", "D", 15000, "Fixture PT One", "IKUTI_PT"])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    assert harness.post("/admin/pegawai/import", {"file": (output, "import.xlsx")}).status_code == 302
    assert not harness.query("SELECT id FROM pegawai WHERE id_pegawai='FORGED1'")


def test_system_settings_separate_versioned_saves_and_retired_routes(harness):
    harness.login("super@fixture.test")
    assert harness.client.get("/superadmin/system-settings").status_code == 200
    assert harness.post("/superadmin/system-settings", {"section": "ppn", "settings_version": "1"}).status_code == 302
    with harness.app.app_context():
        assert not web.get_ppn_enabled()
    assert harness.post("/superadmin/system-settings", {"section": "products", "products": "reg", "settings_version": "1"}).status_code == 409
    assert harness.post("/superadmin/system-settings", {"section": "products", "products": "reg", "settings_version": "2"}).status_code == 302
    with harness.app.app_context():
        assert not web.get_ppn_enabled()
        assert web.get_enabled_products() == ["reg"]
    assert harness.post("/superadmin/products", {"products": "urg", "ppn_enabled": "1"}).status_code == 303
    assert harness.post("/superadmin/settings", {"form_type": "runtime_force_limit", "runtime_force_limit": "1", "force_confirmed": "1"}).status_code == 303
    assert not harness.query("SELECT * FROM app_settings WHERE `key`='runtime_force_limit'")
    html = harness.client.get("/superadmin/settings").get_data(as_text=True)
    assert 'name="ppn_enabled"' not in html and 'name="runtime_force_limit"' not in html


def test_filtered_export_matches_history_and_tenant_scope(harness):
    harness.query("UPDATE transactions SET keterangan='chosen-marker' WHERE id=1")
    for email, expected in (("super@fixture.test", {1}), ("admin@example.com", {1}), ("admin2@fixture.test", set())):
        harness.login(email)
        params = {"start": regression.TODAY, "end": regression.TODAY, "q": "Employee Fixture 1", "product": "reg", "status": "on-proses", "project": "Project 1", "company": "Fixture PT One"}
        contexts = []
        def capture(sender, template, context, **extra):
            contexts.append(context)
        with template_rendered.connected_to(capture, harness.app):
            response = harness.client.get("/superadmin/riwayat" if email.startswith("super") else "/admin/riwayat", query_string=params)
        assert response.status_code == 200
        assert {row["id"] for row in contexts[-1]["rows"]} == expected
        for file_format in ("csv", "xlsx"):
            response = harness.client.get("/admin/export_filtered", query_string={**params, "format": file_format})
            assert {int(row["ID"]) for row in harness.export_rows(response, file_format)} == expected
    harness.login("employee1@fixture.test")
    contexts = []
    with template_rendered.connected_to(capture, harness.app):
        assert harness.client.get("/riwayat").status_code == 200
    assert {row["id"] for row in contexts[-1]["rows"]} == {1, 11}


def test_effective_cycle_counts_do_not_double_count_weekly(harness):
    harness.query("UPDATE clients SET cutoff_mingguan_aktif=1 WHERE id=1")
    with harness.app.app_context():
        assert web.monthly_cycle_counts(get_db()) == (1, 0, 0, 0)
        assert web.get_weekly_cutoff_counts(get_db()) == {2: 2}
        assert web.monthly_cycle_counts(get_db(), "p.perusahaan_induk=?", ["Fixture PT One"]) == (0, 0, 0, 0)


def test_verifier_cli_rejects_mismatch_and_accepts_preserved_policy(harness):
    result = harness.app.test_cli_runner().invoke(args=["verify-company-policy", "--at-date", "2026-12-24", "--days", "14"])
    assert result.exit_code == 0
    assert json.loads(result.output)["checks"] == 28
    harness.query("UPDATE clients SET cutoff_mingguan_aktif=1,cutoff_hari=2 WHERE id=1")
    result = harness.app.test_cli_runner().invoke(args=["verify-company-policy", "--at-date", "2026-12-24"])
    assert result.exit_code != 0 and '"passed": false' in result.output


def test_project_validation_restores_inline_draft(harness):
    harness.login("super@fixture.test")
    response = harness.post("/superadmin/clients/1/projects/add", {"name": "Project 1", "pic_name": "Safe draft"})
    assert response.status_code == 422
    html = response.get_data(as_text=True)
    assert "data-restore-form" in html and "Safe draft" in html
    assert 'data-dialog="project-add"' in html


def test_policy_lock_sql_is_mysql_locking_read_and_fails_closed():
    db = MagicMock()
    lock_policy(db)
    assert "FOR UPDATE" in db.execute.call_args.args[0]
    db.execute.return_value.fetchone.return_value = None
    with pytest.raises(RuntimeError):
        lock_policy(db)


def test_cutoff_write_and_withdrawal_are_serialized(harness):
    harness.prepare_company_limit_withdrawal()
    superadmin = harness.app.test_client()
    harness.login("super@fixture.test", superadmin)
    data = harness.withdrawal_data()
    entered = threading.Event()
    release = threading.Event()
    original = clients.guard_cutoff_change
    def paused_guard(*args, **kwargs):
        entered.set()
        assert release.wait(10)
        return original(*args, **kwargs)
    with patch.object(clients, "guard_cutoff_change", paused_guard):
        with ThreadPoolExecutor(max_workers=2) as executor:
            policy = executor.submit(save_policy, harness, "cutoff", superadmin, cutoff_mingguan_aktif="1", cutoff_hari="2")
            assert entered.wait(10)
            withdrawal = executor.submit(harness.post, "/tarik-gaji", {**data, "nominal": "1000"})
            time.sleep(0.1)
            assert not withdrawal.done()
            release.set()
            assert policy.result(timeout=15).status_code == 302
            assert withdrawal.result(timeout=15).status_code == 302
    saved = harness.query("SELECT periode,nominal FROM transactions")
    assert len(saved) == 1 and saved[0]["nominal"] == 1000
    assert saved[0]["periode"].startswith("W:")


def filter_page(harness, path, **params):
    contexts = []
    def capture(sender, template, context, **extra):
        contexts.append(context)
    with template_rendered.connected_to(capture, harness.app):
        response = harness.client.get(path, query_string=params) if params else harness.client.get(path)
    assert response.status_code == 200
    return response, contexts[-1]

@pytest.mark.parametrize("params,expected", [
    ({}, {1, 2, 3}),
    ({"client": "Fixture PT One"}, {1, 3}),
    ({"client": "Fixture PT Two"}, {2}),
    ({"client": "  fixture pt ONE  "}, {1, 3}),
    ({"client": "Unknown PT"}, set()),
    ({"client": "Fixture PT One", "company": "Project 1"}, {1, 3}),
    ({"client": "Fixture PT One", "company": "Project 2"}, set()),
    ({"client": "Fixture PT Two", "q": "E2"}, {2}),
    ({"client": "Fixture PT Two", "q": "E1"}, set()),
])
def test_employee_company_filter_applies_to_rows_and_projects(harness, params, expected):
    harness.query("UPDATE pegawai SET status_aktif=0 WHERE id=3")
    harness.login("super@fixture.test")
    response, context = filter_page(harness, "/admin/pegawai", **params)
    assert {row["id"] for row in context["rows"]} == expected
    assert context["total"] == len(expected)
    company = params.get("client", "").strip().lower()
    expected_projects = {"Project 1", "Project 2"} if not company else {
        "fixture pt one": {"Project 1"}, "fixture pt two": {"Project 2"}}.get(company, set())
    assert set(context["companies"]) == expected_projects
    if company == "fixture pt one":
        assert context["selected_client"] == "Fixture PT One"
    assert "perusahaan induk" not in response.get_data(as_text=True).lower()
    if params.get("company") == "Project 2" and company == "fixture pt one":
        assert 'value="Project 2" selected' in response.get_data(as_text=True)

@pytest.mark.parametrize("email,expected,projects", [
    ("admin@example.com", {1, 3}, {"Project 1"}),
    ("admin2@fixture.test", {2}, {"Project 2"}),
    ("empty@fixture.test", set(), set()),
])
def test_employee_filters_cannot_expand_admin_scope(harness, email, expected, projects):
    harness.login(email)
    response, context = filter_page(harness, "/admin/pegawai", client="Fixture PT Two")
    assert {row["id"] for row in context["rows"]} == expected
    assert set(context["companies"]) == projects
    assert all(set(entry["projects"]) <= projects for entry in context["project_filter_catalog"])
    assert "perusahaan induk" not in response.get_data(as_text=True).lower()

def test_legacy_project_admin_has_consistent_table_and_filter_scope(harness):
    harness.query("UPDATE admins SET company='Project 1' WHERE id=2")
    harness.login()
    _, context = filter_page(harness, "/admin/pegawai")
    assert {row["id"] for row in context["rows"]} == {1, 3}
    assert context["companies"] == ["Project 1"]
    assert len(context["project_filter_catalog"]) == 1

@pytest.mark.parametrize("path", ["/admin/riwayat", "/superadmin/riwayat"])
def test_history_company_and_project_filters_stay_consistent(harness, path):
    harness.login("super@fixture.test")
    params = {"start": regression.TODAY, "end": regression.TODAY}
    _, context = filter_page(harness, path, **params, company="  FIXTURE PT ONE ")
    assert {row["id"] for row in context["rows"]} == {1, 11}
    assert context["f_company"] == "Fixture PT One"
    assert context["projects"] == ["Project 1"]
    assert {entry["company"] for entry in context["project_filter_catalog"]} == {"Fixture PT One", "Fixture PT Two"}
    response, context = filter_page(harness, path, **params, company="Fixture PT One", project="Project 2")
    assert context["rows"] == []
    assert 'value="Project 2" selected' in response.get_data(as_text=True)
    _, context = filter_page(harness, path, **params, company="Unknown PT")
    assert context["rows"] == [] and context["projects"] == []

def test_history_project_catalog_never_leaks_other_admin_tenants(harness):
    harness.login()
    _, context = filter_page(harness, "/admin/riwayat", company="Fixture PT Two", start=regression.TODAY, end=regression.TODAY)
    assert {row["id"] for row in context["rows"]} == {1, 11}
    assert context["projects"] == ["Project 1"]
    assert context["project_filter_catalog"] == [{"company": "Fixture PT One", "projects": ["Project 1"]}]
    assert context["f_company"] == "Fixture PT One"

@pytest.mark.parametrize("path", ["/superadmin/dashboard", "/admin/dashboard"])
def test_dashboard_company_scopes_metrics_charts_fees_and_cache_not_approval(harness, path):
    harness.query("UPDATE transactions SET status='sukses',nominal=12000,admin_fee=300 WHERE id=1")
    harness.query("UPDATE transactions SET status='sukses',nominal=26000,admin_fee=700 WHERE id=2")
    harness.query("UPDATE pegawai SET status_aktif=0 WHERE id=3")
    harness.query("UPDATE clients SET cutoff_mingguan_aktif=1,cutoff_hari=2 WHERE id=1")
    harness.login("super@fixture.test")
    cases = [
        ("", 3, 2, 38000, 1000, 2, {2: 1}, 1),
        ("  FIXTURE PT ONE ", 2, 1, 12000, 300, 1, {2: 1}, 0),
        ("Fixture PT Two", 1, 1, 26000, 700, 1, {}, 1),
        ("Unknown PT", 0, 0, 0, 0, 0, {}, 0),
        ("Fixture PT One", 2, 1, 12000, 300, 1, {2: 1}, 0),
        ("", 3, 2, 38000, 1000, 2, {2: 1}, 1),
    ]
    for company, employees, registered, nominal, fee, pending, weekly, cycle_a in cases:
        _, context = filter_page(harness, path, company=company, admin_fee_month=regression.PERIOD)
        assert context["total_pegawai"] == employees
        assert context["total_register"] == registered
        assert context["reg_aktif"] == registered
        assert context["eligible"] == registered
        assert context["trx_count"] == registered * 2
        assert context["trx_sum"] == nominal
        assert context["admin_fee_reg_total"] == fee
        assert context["admin_fee_urg_total"] == 0
        assert context["total_admin"] == fee
        assert context["admin_fee_recap_total"] == fee
        assert context["admin_fee_recap_count"] == registered
        assert context["pending_count"] == pending
        assert context["not_borrowed"] == 0
        assert context["inactive_count"] == int(employees > registered)
        assert context["weekly_cutoff_counts"] == weekly
        assert context["cycle_a"] == cycle_a
        assert sum(context["chart_values"]) == nominal
        assert sum(context["trend_values"]) == nominal
        assert sum(row["total"] for row in context["cohort_rows"]) == 0
        queues = context["pending_reg"] + context["pending_urg"] + context["pending_transfer"]
        assert {row["id"] for row in queues} == {11, 12}
        assert {row["nama"] for row in context["recent"]} == {"Employee Fixture 1", "Employee Fixture 2"}
        if company.strip().lower() == "fixture pt one":
            assert context["company_selected"] == "Fixture PT One"

def test_admin_company_filter_normalizes_names_without_widening_unknown(harness):
    harness.login("super@fixture.test")
    _, context = filter_page(harness, "/superadmin/admins", company="  fixture pt ONE ")
    assert {row["id"] for row in context["rows"]} == {2, 4}
    assert context["f_company"] == "Fixture PT One"
    _, context = filter_page(harness, "/superadmin/admins", company="Unknown PT")
    assert context["rows"] == []


def test_company_override_panel_link_opens_filtered_employees(harness):
    from urllib.parse import urlsplit, parse_qs
    from html import unescape
    import re
    harness.query("UPDATE pegawai SET reg_withdrawal_max_override=500000 WHERE id=1")
    harness.login("super@fixture.test")
    html = harness.client.get("/superadmin/clients/1").get_data(as_text=True)
    links = [unescape(link) for link in re.findall(r'href="([^"]+)"', html)]
    target = next(link for link in links if urlsplit(link).path == "/admin/pegawai" and
                  parse_qs(urlsplit(link).query).get("client") == ["Fixture PT One"])
    _, context = filter_page(harness, target)
    assert {row["id"] for row in context["rows"]} == {1, 3}

def test_dashboard_cohorts_and_trends_do_not_mix_companies(harness):
    today = date.fromisoformat(regression.TODAY)
    previous = today.replace(day=1) - timedelta(days=1)
    earlier = previous.replace(day=1) - timedelta(days=1)
    harness.query("INSERT INTO transactions(user_id,tanggal,periode,nominal,admin_fee,status,created_at,product) VALUES(1,?,?,11000,100,'sukses',?,'reg')",
                  (earlier.isoformat(), earlier.strftime("%Y-%m"), earlier.isoformat()))
    harness.query("INSERT INTO transactions(user_id,tanggal,periode,nominal,admin_fee,status,created_at,product) VALUES(1,?,?,13000,100,'sukses',?,'reg')",
                  (previous.isoformat(), previous.strftime("%Y-%m"), previous.isoformat()))
    harness.query("INSERT INTO transactions(user_id,tanggal,periode,nominal,admin_fee,status,created_at,product) VALUES(2,?,?,17000,100,'sukses',?,'reg')",
                  (earlier.isoformat(), earlier.strftime("%Y-%m"), earlier.isoformat()))
    harness.login("super@fixture.test")
    for company, total, repeat, amount in [("Fixture PT One", 1, 1, 24000), ("Fixture PT Two", 1, 0, 17000), ("", 2, 1, 41000), ("Unknown PT", 0, 0, 0)]:
        _, context = filter_page(harness, "/superadmin/dashboard", company=company)
        cohort = next(row for row in context["cohort_rows"] if row["cohort"] == earlier.strftime("%Y-%m"))
        assert (cohort["total"], cohort["repeat"]) == (total, repeat)
        assert sum(context["trend_values"]) == amount

@pytest.mark.parametrize("company,expected", [("Fixture PT One", {"Project 1", "Legacy Project"}), ("Fixture PT Two", {"Project 2"})])
def test_project_options_include_inactive_legacy_rows_without_other_tenants(harness, company, expected):
    harness.query("UPDATE pegawai SET status_aktif=0,perusahaan='Legacy Project' WHERE id=3")
    harness.login("super@fixture.test")
    _, context = filter_page(harness, "/admin/pegawai", client=company)
    assert set(context["companies"]) == expected
