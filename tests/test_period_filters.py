import csv
import io
import os
from datetime import date, datetime

import pytest
from flask import template_rendered
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

os.environ["START_INLINE_WORKER"] = "0"
os.environ["EXPORT_CSV_ENABLED"] = "0"

from app import create_app
from app.db import get_db
from app.routes import web
from app.services import company_lifecycle, security_service
from app.utils.cache import _cache

PASSWORD = "FixturePeriodPassword123"
PASSWORD_HASH = generate_password_hash(PASSWORD)
TODAY = date(2026, 10, 19)
CURRENT_MARKERS = {"month-start", "current-month", "month-end"}


class PeriodHarness:
    def __init__(self, app):
        self.app = app
        self.client = app.test_client()

    def query(self, sql, params=()):
        with self.app.app_context():
            connection = get_db()
            cursor = connection.execute(sql, params)
            rows = [dict(row) for row in cursor.fetchall()] if cursor.description else []
            connection.commit()
            return rows

    def login(self, email):
        self.client.get("/login")
        with self.client.session_transaction() as session:
            token = session["_csrf_token"]
        response = self.client.post("/login", data={"email": email, "password": PASSWORD, "csrf_token": token})
        assert response.status_code == 302

    def render(self, path, params=None):
        contexts = []

        def capture(sender, template, context, **extra):
            contexts.append(context)

        with template_rendered.connected_to(capture, self.app):
            response = self.client.get(path, query_string=params)
        assert response.status_code == 200
        assert contexts
        return response.get_data(as_text=True), contexts[-1]


@pytest.fixture
def period_harness(tmp_path, monkeypatch):
    _cache.clear()
    monkeypatch.setattr(web, "business_today", lambda: TODAY)
    monkeypatch.setattr(company_lifecycle, "business_now", lambda: datetime(2026, 10, 19, 12, tzinfo=company_lifecycle.WIB))
    monkeypatch.setattr(web, "enqueue_email", lambda *args, **kwargs: None)
    app = create_app({"TESTING": True, "DB_TYPE": "sqlite", "DB_PATH": str(tmp_path / "periods.db"),
                      "APP_SECRET": "isolated-period-fixture-not-production", "SESSION_COOKIE_SECURE": False,
                      "ADMIN_PASSWORD": "", "SUPERADMIN_PASSWORD": "", "SEED_DEMO_DATA": False,
                      "ALLOW_DATE_OVERRIDE": False, "ALLOW_DESTRUCTIVE_RESET": False})
    harness = PeriodHarness(app)
    with app.app_context():
        connection = get_db()
        connection.execute("DELETE FROM client_projects")
        connection.execute("DELETE FROM clients")
        connection.execute("DELETE FROM admins")
        for client_id, company in ((1, "Fixture PT One"), (2, "Fixture PT Two")):
            connection.execute("INSERT INTO clients (id,name,name_key,allow_projects,created_at,updated_at) VALUES (?,?,?,1,?,?)",
                               (client_id, company, company.lower(), TODAY.isoformat(), TODAY.isoformat()))
        for admin_id, email, company, role in ((1, "super@fixture.test", None, "superadmin"),
                                               (2, "admin@fixture.test", "Fixture PT One", "admin")):
            connection.execute("INSERT INTO admins (id,name,email,password_hash,company,role,status_aktif,hak_approval,created_at) "
                               "VALUES (?,?,?,?,?,?,1,1,?)", (admin_id, email, email, PASSWORD_HASH, company, role, TODAY.isoformat()))
        for employee_id, company in ((1, "Fixture PT One"), (2, "Fixture PT Two")):
            email = f"employee{employee_id}@fixture.test"
            connection.execute("INSERT INTO pegawai (id,id_pegawai,nama,email,gaji,status_aktif,perusahaan,perusahaan_induk,no_rekening,siklus_gaji,created_at) "
                               "VALUES (?,?,?,?,6000000,1,?,?,?,'A',?)",
                               (employee_id, f"E{employee_id}", f"Employee {employee_id}", email, f"Project {employee_id}", company, "123456", TODAY.isoformat()))
            connection.execute("INSERT INTO user_accounts (pegawai_id,name,email,password_hash,status_aktif,created_at) VALUES (?,?,?,?,1,?)",
                               (employee_id, f"Employee {employee_id}", email, PASSWORD_HASH, TODAY.isoformat()))
            connection.execute("INSERT INTO users (id,name,email,gaji,created_at) VALUES (?,?,?,6000000,?)",
                               (employee_id, f"Employee {employee_id}", email, TODAY.isoformat()))
        for transaction_id, employee_id, day, period, marker in (
            (101, 1, "2026-09-30", "2026-09", "previous-month"),
            (102, 1, "2026-10-01", "2026-09", "month-start"),
            (103, 1, "2026-10-19", "2026-10", "current-month"),
            (104, 1, "2026-10-31", "W:2026-10-29", "month-end"),
            (105, 1, "2026-11-01", "2026-11", "next-month"),
            (106, 1, "2026-05-31", "2026-05", "older-month"),
            (201, 2, "2026-10-19", "2026-10", "other-company"),
        ):
            connection.execute("INSERT INTO transactions (id,user_id,tanggal,periode,nominal,admin_fee,status,created_at,product,keterangan) "
                               "VALUES (?,?,?,?,1000,100,'sukses',?,'reg',?)",
                               (transaction_id, employee_id, day, period, day + "T12:00:00", marker))
        connection.commit()
    yield harness
    _cache.clear()


@pytest.mark.parametrize(("email", "path", "extra_markers"), [
    ("admin@fixture.test", "/admin/riwayat", set()),
    ("super@fixture.test", "/superadmin/riwayat", {"other-company"}),
])
@pytest.mark.parametrize("params", [None, {"start": "", "end": ""}, {"start": "bad-date", "end": "2026-10-19"},
                                     {"start": "2026-02-30", "end": "2026-10-19"}])
def test_management_history_defaults_to_current_month(period_harness, email, path, extra_markers, params):
    period_harness.login(email)
    html, context = period_harness.render(path, params)
    assert context["start"] == "2026-10-01"
    assert context["end"] == "2026-10-31"
    assert {row["keterangan"] for row in context["rows"]} == CURRENT_MARKERS | extra_markers
    assert 'id="history-start" value="2026-10-01"' in html
    assert 'id="history-end" value="2026-10-31"' in html
    assert 'type="submit"' in html and "Terapkan" in html


@pytest.mark.parametrize(("email", "path"), [("admin@fixture.test", "/admin/riwayat"),
                                             ("super@fixture.test", "/superadmin/riwayat")])
def test_management_history_honors_applied_dates_and_reset(period_harness, email, path):
    period_harness.login(email)
    for start, end in (("2026-09-01", "2026-09-30"), ("2026-09-30", "2026-09-01")):
        _, context = period_harness.render(path, {"start": start, "end": end})
        assert context["start"] == "2026-09-01"
        assert context["end"] == "2026-09-30"
        assert {row["keterangan"] for row in context["rows"]} == {"previous-month"}
    _, context = period_harness.render(path)
    assert context["start"] == "2026-10-01"
    assert context["end"] == "2026-10-31"


@pytest.mark.parametrize("period", [None, "", "invalid", "2026-13", "2026-00", "2026-1", "all"])
def test_fee_recap_defaults_to_current_month(period_harness, period):
    period_harness.login("super@fixture.test")
    html, context = period_harness.render("/superadmin/dashboard", None if period is None else {"admin_fee_month": period})
    assert context["admin_fee_month"] == "2026-10"
    assert context["admin_fee_recap_total"] == 400
    assert context["admin_fee_recap_count"] == 4
    assert 'id="admin_fee_month" value="2026-10" required' in html
    assert ">Bulan ini</a>" in html


def test_fee_recap_honors_selected_month_even_with_warm_cache(period_harness):
    period_harness.login("super@fixture.test")
    period_harness.render("/superadmin/dashboard")
    _, context = period_harness.render("/superadmin/dashboard", {"admin_fee_month": "2026-09"})
    assert context["admin_fee_month"] == "2026-09"
    assert context["admin_fee_recap_total"] == 100
    _, context = period_harness.render("/superadmin/dashboard")
    assert context["admin_fee_recap_total"] == 400


def test_dashboard_export_dates_follow_fee_recap_month(period_harness):
    period_harness.login('super@fixture.test')
    html, context = period_harness.render('/superadmin/dashboard', {'admin_fee_month': '2026-09'})
    assert context['export_start'] == '2026-09-01'
    assert context['export_end'] == '2026-09-30'
    assert 'id="export-start" name="start" value="2026-09-01"' in html
    assert 'id="export-end" name="end" value="2026-09-30"' in html
    assert 'Semua perusahaan' in html


def test_dashboard_export_company_options_are_parent_companies(period_harness):
    period_harness.login('super@fixture.test')
    html, context = period_harness.render('/superadmin/dashboard')
    assert set(context['companies']) == {'Fixture PT One', 'Fixture PT Two'}
    assert '<option value="Fixture PT One"' in html
    assert '<option value="Project 1"' not in html


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
@pytest.mark.parametrize("ppn_enabled", ["0", "1"])
@pytest.mark.parametrize(("month", "expected_total", "expected_products"), [
    ("2026-09", 12000, {"reg": 12000}),
    ("2026-10", 64381, {"reg": 19981, "urg": 44400}),
    ("2026-11", 22222, {"reg": 22222}),
])
def test_success_fees_match_dashboard_and_exports(period_harness, file_format, ppn_enabled,
                                                 month, expected_total, expected_products):
    for transaction_id, product, fee in ((101, 'reg', 12000), (102, 'reg', 16650),
                                         (103, 'urg', 27750), (104, 'reg', 3331),
                                         (201, 'urg', 16650), (105, 'reg', 22222)):
        period_harness.query('UPDATE transactions SET product=?, admin_fee=? WHERE id=?',
                             (product, fee, transaction_id))
    for transaction_id, status, product, fee in ((301, 'on-proses', 'reg', 999999),
                                                 (302, 'ditolak', 'urg', 888888),
                                                 (303, 'dibatalkan', 'reg', 777777)):
        period_harness.query(
            'INSERT INTO transactions (id,user_id,tanggal,periode,nominal,admin_fee,status,created_at,product) '
            'VALUES (?,1,?,?,?,?,?,?,?)',
            (transaction_id, '2026-10-19', '2026-09', 1000, fee, status, '2026-10-19T12:00:00', product))
    period_harness.query('INSERT OR REPLACE INTO app_settings (key,value) VALUES (?,?)',
                         ('ppn_enabled', ppn_enabled))
    period_harness.query('UPDATE pegawai SET admin_fee_flat=999999')
    period_harness.login('super@fixture.test')
    start, end = web.month_filter_bounds(date.fromisoformat(month + '-01'))
    _, context = period_harness.render('/superadmin/dashboard', {'admin_fee_month': month})
    _, warm_context = period_harness.render('/superadmin/dashboard', {'admin_fee_month': month})
    assert context['admin_fee_recap_total'] == expected_total
    assert warm_context['admin_fee_recap_total'] == expected_total
    if month == '2026-10':
        assert context['total_admin'] == expected_total
        assert context['admin_fee_reg_total'] == expected_products['reg']
        assert context['admin_fee_urg_total'] == expected_products['urg']
    dashboard_products = {row['product']: int(row['admin_fee_total'])
                          for row in context['admin_fee_recap_rows']}
    assert dashboard_products == expected_products
    for endpoint, query in (
        ('/superadmin/export_range', {'start': start.isoformat(), 'end': end.isoformat()}),
        ('/admin/export', {'periode': month}),
    ):
        response = period_harness.client.get(endpoint, query_string={**query, 'format': file_format})
        assert response.status_code == 200
        if file_format == 'xlsx':
            workbook = load_workbook(io.BytesIO(response.data))
            values = list(workbook.active.values)
            rows = [dict(zip(values[0], row)) for row in values[1:]]
            workbook.close()
        else:
            rows = list(csv.DictReader(io.StringIO(response.data.decode('utf-8-sig'))))
        assert len({row['ID'] for row in rows}) == len(rows)
        successful = [row for row in rows if row['Status'] == 'sukses']
        assert sum(int(row['Admin']) for row in successful) == expected_total
        assert len(successful) == context['admin_fee_recap_count']
        export_products = {}
        for row in successful:
            export_products[row['Produk']] = export_products.get(row['Produk'], 0) + int(row['Admin'])
        assert export_products == dashboard_products


@pytest.mark.parametrize('file_format', ['csv', 'xlsx'])
def test_success_fee_card_includes_legacy_regular_product(period_harness, file_format):
    period_harness.query('UPDATE transactions SET product=NULL, admin_fee=12345 WHERE id=103')
    period_harness.login('super@fixture.test')
    _, context = period_harness.render('/superadmin/dashboard')
    response = period_harness.client.get('/superadmin/export_range', query_string={
        'start': '2026-10-01', 'end': '2026-10-31', 'format': file_format,
    })
    if file_format == 'xlsx':
        workbook = load_workbook(io.BytesIO(response.data))
        values = list(workbook.active.values)
        rows = [dict(zip(values[0], row)) for row in values[1:]]
        workbook.close()
    else:
        rows = list(csv.DictReader(io.StringIO(response.data.decode('utf-8-sig'))))
    assert sum(int(row['Admin']) for row in rows if row['Status'] == 'sukses') == 12645
    assert context['total_admin'] == context['admin_fee_recap_total'] == 12645
    assert context['admin_fee_reg_total'] == 12645


def test_success_fee_reconciles_after_approval_with_warm_cache(period_harness):
    period_harness.query("UPDATE transactions SET status='on-proses', product='urg', admin_fee=16650, "
                         "admin_approved_at=?, admin_approved_by='Fixture Admin' WHERE id=103",
                         ('2026-10-19T11:00:00',))
    period_harness.login('super@fixture.test')
    _, before = period_harness.render('/superadmin/dashboard')
    assert before['total_admin'] == before['admin_fee_recap_total'] == 300
    with period_harness.client.session_transaction() as session:
        token = session['_csrf_token']
    response = period_harness.client.post('/superadmin/tx/103/approve', data={'csrf_token': token})
    assert response.status_code == 302
    _, after = period_harness.render('/superadmin/dashboard')
    assert after['total_admin'] == after['admin_fee_recap_total'] == 16950
    response = period_harness.client.get('/superadmin/export_range', query_string={
        'start': '2026-10-01', 'end': '2026-10-31', 'format': 'csv',
    })
    rows = list(csv.DictReader(io.StringIO(response.data.decode('utf-8-sig'))))
    assert sum(int(row['Admin']) for row in rows if row['Status'] == 'sukses') == after['total_admin']


def test_scheduled_excel_success_fees_match_dashboard(period_harness, monkeypatch):
    from unittest.mock import mock_open, patch
    from app.tasks import export_scheduler

    period_harness.query("UPDATE transactions SET status='on-proses', admin_fee=900000 WHERE id=103")
    period_harness.login('super@fixture.test')
    _, context = period_harness.render('/superadmin/dashboard')
    monkeypatch.setattr(export_scheduler, 'business_today', lambda: TODAY)
    with period_harness.app.app_context(), \
         patch.object(export_scheduler, 'open', mock_open(), create=True), \
         patch.object(export_scheduler.os, 'makedirs'):
        result = export_scheduler._export_current_period()
    workbook = load_workbook(io.BytesIO(result[2]))
    values = list(workbook.active.values)
    rows = [dict(zip(values[0], row)) for row in values[1:]]
    workbook.close()
    assert sum(row['Admin'] for row in rows if row['Status'] == 'sukses') == 300
    assert context['total_admin'] == context['admin_fee_recap_total'] == 300


@pytest.mark.parametrize("period", [None, "", "invalid", "2026-13", "2026-00"])
def test_employee_history_defaults_to_calendar_month(period_harness, period):
    period_harness.login("employee1@fixture.test")
    html, context = period_harness.render("/riwayat", None if period is None else {"periode": period})
    assert context["periode"] == "2026-10"
    assert {row["keterangan"] for row in context["rows"]} == CURRENT_MARKERS
    assert 'value="2026-10" selected' in html
    assert 'value="last-6" selected' not in html


def test_employee_other_periods_require_explicit_selection(period_harness):
    period_harness.login("employee1@fixture.test")
    _, context = period_harness.render("/riwayat", {"periode": "2026-09"})
    assert context["periode"] == "2026-09"
    assert {row["keterangan"] for row in context["rows"]} == {"previous-month"}
    _, context = period_harness.render("/riwayat", {"periode": "last-6"})
    assert context["periode"] == "last-6"
    assert {row["keterangan"] for row in context["rows"]} == {"previous-month", "month-start", "current-month", "older-month"}
    _, context = period_harness.render("/riwayat")
    assert context["periode"] == "2026-10"


def test_weekly_employee_still_defaults_to_calendar_month(period_harness):
    period_harness.query("UPDATE admins SET cutoff_mingguan_aktif=1, cutoff_hari=2 WHERE id=2")
    for transaction in period_harness.query("SELECT id,tanggal FROM transactions WHERE user_id=1"):
        period = web.weekly_period_key(date.fromisoformat(transaction["tanggal"]), 2)
        period_harness.query("UPDATE transactions SET periode=? WHERE id=?", (period, transaction["id"]))
    period_harness.login("employee1@fixture.test")
    _, context = period_harness.render("/riwayat")
    assert context["periode"] == "2026-10"
    assert {row["keterangan"] for row in context["rows"]} == CURRENT_MARKERS
    weekly_period = web.weekly_period_key(TODAY, 2)
    _, context = period_harness.render("/riwayat", {"periode": weekly_period})
    assert context["periode"] == weekly_period
    assert {row["keterangan"] for row in context["rows"]} == {"current-month"}
    _, context = period_harness.render("/riwayat", {"periode": "last-6"})
    assert context["periode"] == "last-6"
    assert "previous-month" in {row["keterangan"] for row in context["rows"]}


@pytest.mark.parametrize(("email", "path"), [("admin@fixture.test", "/admin/dashboard"),
                                             ("admin@fixture.test", "/admin/riwayat"),
                                             ("super@fixture.test", "/superadmin/dashboard"),
                                             ("super@fixture.test", "/superadmin/riwayat")])
def test_export_forms_start_at_current_month(period_harness, email, path):
    period_harness.login(email)
    html, context = period_harness.render(path)
    assert context["filter_current_month"] == "2026-10"
    if email.startswith("super"):
        assert 'id="export-start" name="start" value="2026-10-01"' in html
        assert 'id="export-end" name="end" value="2026-10-31"' in html
    else:
        assert 'list="export-period-options" value="2026-10"' in html
        assert '<option value="all">Semua periode</option>' in html


@pytest.mark.parametrize(("period", "expected_ids"), [
    (None, {102, 103, 104}), ("", {102, 103, 104}), ("invalid", {102, 103, 104}),
    ("2026-13", {102, 103, 104}), ("2026-09", {101}),
    ("W:2026-10-29", {104}), ("all", {101, 102, 103, 104, 105, 106}),
])
@pytest.mark.parametrize("file_format", [None, "csv", "xlsx"])
def test_admin_export_default_and_explicit_periods(period_harness, period, expected_ids, file_format):
    period_harness.login("admin@fixture.test")
    query = {} if period is None else {"periode": period}
    if file_format:
        query["format"] = file_format
    response = period_harness.client.get("/admin/export", query_string=query)
    assert response.status_code == 200
    if file_format == "xlsx":
        workbook = load_workbook(io.BytesIO(response.data))
        rows = list(workbook.active.values)
        workbook.close()
    else:
        rows = list(csv.reader(io.StringIO(response.get_data(as_text=True).lstrip("\ufeff"))))
    assert f'.{file_format or "csv"}' in response.headers['Content-Disposition']
    assert {int(row[0]) for row in rows[1:]} == expected_ids


@pytest.mark.parametrize(("instant", "month", "start", "end"), [
    ("2026-10-31T16:59:59+00:00", "2026-10", "2026-10-01", "2026-10-31"),
    ("2026-10-31T17:00:00+00:00", "2026-11", "2026-11-01", "2026-11-30"),
    ("2026-12-31T17:00:00+00:00", "2027-01", "2027-01-01", "2027-01-31"),
    ("2028-02-28T17:00:00+00:00", "2028-02", "2028-02-01", "2028-02-29"),
])
def test_defaults_follow_real_jakarta_month_boundary(period_harness, monkeypatch, instant, month, start, end):
    frozen = datetime.fromisoformat(instant)

    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return frozen.astimezone(tz)

    monkeypatch.setattr(security_service, "datetime", FrozenDateTime)
    monkeypatch.setattr(web, "business_today", security_service.business_today)
    period_harness.login("super@fixture.test")
    _, context = period_harness.render("/superadmin/riwayat")
    assert context["start"] == start
    assert context["end"] == end
    assert context["filter_current_month"] == month


def test_history_defaults_do_not_restore_simulation_or_old_selection(period_harness):
    period_harness.app.config["ALLOW_DATE_OVERRIDE"] = True
    period_harness.login("employee1@fixture.test")
    with period_harness.client.session_transaction() as session:
        session["sim_date"] = "2025-01-15"
    _, context = period_harness.render("/riwayat", {"tanggal": "2025-02-15"})
    assert context["periode"] == "2026-10"
    assert {row["keterangan"] for row in context["rows"]} == CURRENT_MARKERS
    assert web.period_key_by_cycle(date(2026, 10, 1), "B") == "2026-09"
