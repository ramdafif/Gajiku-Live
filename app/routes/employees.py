import io
import json
import secrets
import sqlite3
import time
import zipfile

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from defusedxml.common import DefusedXmlException

from app.db import get_db
from app.services import employee_changes as changes
from app.services.employee_finance import EMPLOYEE_FIELDS, FINANCIAL_FIELDS, financial_values
from app.services.client_service import client_by_name, list_clients, list_projects, name_key
from app.services.security_service import rate_limit

bp = Blueprint('employees', __name__)


@bp.before_request
def staff_only():
    from app.routes.web import require_admin
    return require_admin()


@bp.post('/superadmin/clients/<int:client_id>/employee-policy')
def company_policy(client_id):
    if not session.get('is_superadmin'):
        abort(403)
    from app.services.employee_finance import configure_employee_policy
    try:
        with changes.write_transaction([client_id]) as db:
            row = db.execute('SELECT * FROM clients WHERE id=?', (client_id,)).fetchone()
            if not row:
                abort(404)
            configure_employee_policy(db, dict(row), request.form.get('plafon_start_day'),
                                      request.form.get('default_admin_fee'))
        flash('Fee PT tersimpan. Perubahan tanggal periode existing berlaku setelah periode berjalan selesai.', 'success')
    except ValueError as error:
        flash(str(error), 'error')
    return redirect(url_for('clients.superadmin_client_detail', client_id=client_id))


def form_values(form, existing=None):
    values = {field: form[field] for field in EMPLOYEE_FIELDS if field in form}
    if not session.get('is_superadmin') and 'admin_fee_flat' in form:
        abort(403)
    if 'status' in form:
        values['status_aktif'] = form['status']
    if not existing:
        values.setdefault('status_aktif', 1)
        values.setdefault('perusahaan_induk', session.get('company') or '')
    return values


def announce(db, client, batch_id, pending, immediate=False, bank_changed=False):
    title = 'Menunggu persetujuan pegawai' if pending else 'Pemberitahuan aktivitas pegawai'
    if bank_changed:
        title += ' - Tujuan pencairan berubah'
    targets = db.execute('SELECT employee_id,proposed_json,before_json FROM employee_change_items WHERE batch_id=?', (batch_id,)).fetchall()
    names = []
    for target in targets[:3]:
        employee = db.execute('SELECT nama FROM pegawai WHERE id=?', (target['employee_id'],)).fetchone()
        names.append(employee['nama'] if employee else json.loads(target['proposed_json']).get('nama')
                     or json.loads(target['before_json']).get('nama', 'Pegawai'))
    changes.notify(db, client['id'], 'employee_request' if pending else 'employee_activity', title,
                   f"PT {client['name']}; admin {session.get('admin_email') or changes.actor()}; batch #{batch_id}. "
                   + f"{len(targets)} item: {', '.join(names)}. "
                   + ('Perubahan biasa sudah berlaku. ' if immediate else '')
                   + ('Perubahan terlindungi menunggu persetujuan.' if pending else 'Buka detail aktivitas.'), batch_id)


def save_employee(employee_id=None):
    from app.routes.web import require_admin
    denied = require_admin()
    if denied:
        return denied
    try:
        with changes.write_transaction() as db:
            old = db.execute('SELECT * FROM pegawai WHERE id=?', (employee_id,)).fetchone() if employee_id else None
            if employee_id and not old:
                abort(404)
            if old:
                changes.scoped_client(db, old['perusahaan_induk'])
            incoming = form_values(request.form, old)
            values = changes.validate_employee(db, incoming, old)
            client = changes.scoped_client(db, values['perusahaan_induk'])
            changes.assert_identity(db, values, employee_id)
            if not employee_id and session.get('is_superadmin'):
                changes.create_employee(db, values)
                message = 'Pegawai ditambahkan.'
            else:
                batch_id = changes.new_batch(db, client['id'], 'edit' if employee_id else 'create')
                if employee_id:
                    immediate, pending = changes.submit_update(db, client, employee_id, old, values, batch_id)
                else:
                    changes.queue_item(db, batch_id, 'create', values)
                    changes.audit(db, client['id'], None, 'requested.create', after=values)
                    immediate, pending = False, True
                if not session.get('is_superadmin') and (immediate or pending):
                    bank_changed = bool(old and any(old[field] != values[field] for field in
                                        ('no_rekening', 'no_rekening_lain', 'rekening_ewallet')))
                    announce(db, client, batch_id, pending, immediate, bank_changed)
                message = ('Perubahan terlindungi menunggu persetujuan superadmin. ' if pending else '')
                message += 'Perubahan biasa yang valid langsung berlaku.' if immediate else 'Data diproses.'
        flash(message, 'success')
    except ValueError as error:
        flash(str(error), 'error')
    return redirect(url_for('web.admin_pegawai'))


def delete_employee(employee_id):
    from app.routes.web import require_admin
    denied = require_admin()
    if denied:
        return denied
    try:
        with changes.write_transaction() as db:
            row = db.execute('SELECT * FROM pegawai WHERE id=?', (employee_id,)).fetchone()
            if not row:
                abort(404)
            client = changes.scoped_client(db, row['perusahaan_induk'])
            history = db.execute('''SELECT 1 FROM transactions t JOIN users u ON u.id=t.user_id
                WHERE u.pegawai_id=? OR LOWER(u.email)=LOWER(?) LIMIT 1''', (employee_id, row['email'])).fetchone()
            if history:
                raise ValueError('Pegawai memiliki riwayat transaksi. Nonaktifkan, jangan hapus.')
            changes.cancel_employee_requests(db, employee_id)
            db.execute('CREATE TABLE IF NOT EXISTS pegawai_archive (pegawai_id INTEGER,snapshot TEXT,deleted_at TEXT)')
            db.execute('INSERT INTO pegawai_archive (pegawai_id,snapshot,deleted_at) VALUES (?,?,?)',
                       (employee_id, changes.encoded(dict(row)), changes.lifecycle.timestamp(changes.lifecycle.local_now())))
            db.execute('DELETE FROM user_accounts WHERE pegawai_id=?', (employee_id,))
            db.execute('DELETE FROM users WHERE pegawai_id=?', (employee_id,))
            db.execute('DELETE FROM pegawai WHERE id=?', (employee_id,))
            changes.audit(db, client['id'], employee_id, 'deleted', before=changes.snapshot(row))
            if not session.get('is_superadmin'):
                batch_id = changes.new_batch(db, client['id'], 'delete')
                changes.queue_item(db, batch_id, 'activity', {'deleted': True}, employee_id,
                                   {'nama': row['nama'], 'id_pegawai': row['id_pegawai']})
                db.execute("UPDATE employee_change_items SET status='applied' WHERE batch_id=?", (batch_id,))
                announce(db, client, batch_id, False, True)
        flash('Pegawai diarsipkan dan dihapus; aktivitas tercatat.', 'success')
    except ValueError as error:
        flash(str(error), 'error')
    return redirect(url_for('web.admin_pegawai'))


@bp.get('/employee-changes')
def inbox():
    db = get_db()
    params = []
    scope = ''
    if not session.get('is_superadmin'):
        scope = 'WHERE c.name_key=?'
        params.append(name_key(session.get('company')))
    page = max(1, request.args.get('page', 1, type=int))
    rows = db.execute(f'''SELECT b.*,c.name,
        (SELECT COUNT(*) FROM employee_change_items i WHERE i.batch_id=b.id AND i.status='pending') AS pending_count
        FROM employee_change_batches b JOIN clients c ON c.id=b.client_id {scope}
        ORDER BY b.id DESC LIMIT 50 OFFSET ?''', (*params, (page - 1) * 50)).fetchall()
    return render_template('shared/employee_changes.html', batches=rows, page=page)


def get_batch(db, batch_id):
    row = db.execute('''SELECT b.*,c.name FROM employee_change_batches b
        JOIN clients c ON c.id=b.client_id WHERE b.id=?''', (batch_id,)).fetchone()
    if not row:
        abort(404)
    if not session.get('is_superadmin') and name_key(row['name']) != name_key(session.get('company')):
        abort(404)
    return dict(row)


@bp.route('/employee-changes/<int:batch_id>', methods=['GET', 'POST'])
def detail(batch_id):
    db = get_db()
    batch = get_batch(db, batch_id)
    if request.method == 'POST':
        decision = request.form.get('decision')
        if decision != 'cancel' and not session.get('is_superadmin'):
            abort(403)
        try:
            ids = [int(value) for value in request.form.getlist('item_ids')]
            with changes.write_transaction([batch['client_id']]) as db:
                batch = get_batch(db, batch_id)
                status = changes.decide_items(db, batch, ids, decision, request.form.get('reason', ''))
            flash('Keputusan tersimpan: ' + status, 'success' if status != 'conflict' else 'warning')
        except ValueError as error:
            flash(str(error), 'error')
        return redirect(url_for('employees.detail', batch_id=batch_id))
    items = []
    for row in db.execute('SELECT * FROM employee_change_items WHERE batch_id=? ORDER BY id', (batch_id,)).fetchall():
        item = dict(row)
        item['before'] = json.loads(item['before_json'])
        item['proposed'] = json.loads(item['proposed_json'])
        employee = db.execute('SELECT nama,id_pegawai FROM pegawai WHERE id=?', (item['employee_id'],)).fetchone()
        item['employee_name'] = employee['nama'] if employee else item['proposed'].get('nama') or item['before'].get('nama')
        items.append(item)
    return render_template('shared/employee_change_detail.html', batch=batch, items=items,
                           own_request=batch['actor'] == changes.actor())


IMPORT_FIELDS = [('id_pegawai', 'ID Pegawai*'), ('nama', 'Nama*'), ('email', 'Email*'),
                 ('jabatan', 'Jabatan'), ('no_rekening', 'No Rekening Bank Utama'),
                 ('no_rekening_lain', 'No Rekening Bank Lain'), ('rekening_ewallet', 'Rekening E-Wallet'),
                 ('no_telp', 'No Telp'), ('gaji', 'Gaji Pokok'), ('plafon_mode', 'Mode Plafon*'),
                 ('plafon', 'Plafon Nominal'), ('plafon_persen', 'Plafon Persen'),
                 ('status_aktif', 'Status Aktif'), ('siklus_gaji', 'Siklus Gaji')]


def import_fields(db):
    fields = list(IMPORT_FIELDS)
    if session.get('is_superadmin'):
        fields += [('perusahaan', 'Project'), ('perusahaan_induk', 'Perusahaan*'), ('admin_fee_flat', 'Admin Fee')]
    else:
        client = changes.scoped_client(db, session.get('company'))
        if client['allow_projects']:
            fields.append(('perusahaan', 'Project*'))
    return fields


def import_template():
    from app.routes.web import require_admin
    denied = require_admin()
    if denied:
        return denied
    db = get_db()
    fields = import_fields(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Import Pegawai v2'
    sheet.append([label for field, label in fields])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F2937')
        sheet.column_dimensions[cell.column_letter].width = max(20, len(cell.value) + 2)
    sheet.freeze_panes = 'A2'
    for index, (field, label) in enumerate(fields, 1):
        formula = {'plafon_mode': '"nominal,persentase"', 'status_aktif': '"Aktif,Nonaktif"',
                   'siklus_gaji': '"A,B,C,D"'}.get(field)
        if formula:
            validation = DataValidation(type='list', formula1=formula, allow_blank=field != 'plafon_mode')
            sheet.add_data_validation(validation)
            letter = sheet.cell(1, index).column_letter
            validation.add(f'{letter}2:{letter}501')
    guide = workbook.create_sheet('Petunjuk')
    for text in (
        'Template v2. Maksimal 500 pegawai / 5 MB. Tidak ada baris contoh yang ikut diimpor.',
        'Isi gaji aktual atau biarkan kosong. Jangan masukkan plafon sebagai gaji.',
        'Mode nominal: isi Plafon Nominal. Mode persentase: isi gaji dan Plafon Persen (0-100).',
        'Nominal berupa angka bulat tanpa pemisah ribuan. Rekening dan ID disimpan sebagai teks.',
        'Admin PT: data valid menjadi usulan; pegawai baru belum tersedia sampai disetujui superadmin.',
        'Fee kosong mengikuti default PT. Hanya superadmin dapat menetapkan pengecualian.',
    ):
        guide.append([text])
    guide.column_dimensions['A'].width = 120
    catalog = workbook.create_sheet('Master PT dan Project')
    catalog.append(['Perusahaan', 'Project'])
    clients = list_clients(db) if session.get('is_superadmin') else [changes.scoped_client(db, session.get('company'))]
    for client in clients:
        projects = list_projects(db, client['id']) if client['allow_projects'] else []
        for project in projects or [{'name': ''}]:
            catalog.append([client['name'], project['name']])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='template-pegawai-v2.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def import_employees():
    from app.routes.web import require_admin
    denied = require_admin()
    if denied:
        return denied
    limited = rate_limit('employee-import', changes.actor(), limit=5, window=60)
    if limited:
        return limited
    upload = request.files.get('file')
    if not upload or not upload.filename or not upload.filename.lower().endswith('.xlsx'):
        flash('Pilih file .xlsx yang valid.', 'error')
        return redirect(url_for('web.admin_pegawai'))
    try:
        content = upload.read(5 * 1024 * 1024 + 1)
        if len(content) > 5 * 1024 * 1024:
            raise ValueError('Ukuran maksimal 5 MB.')
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 2000 or sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('File terlalu besar setelah diekstrak.')
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except Exception as error:
            raise ValueError('File tidak bisa dibaca. Pastikan XML Excel valid dan aman.') from error
        try:
            sheet = workbook.active
            if sheet is None or (sheet.max_column and sheet.max_column > 30):
                raise ValueError('Sheet/kolom Excel tidak valid.')
            fields = import_fields(get_db())
            headers = [str(cell.value or '').strip() for cell in next(sheet.iter_rows(max_row=1, max_col=31))]
            while headers and not headers[-1]:
                headers.pop()
            expected = {label: field for field, label in fields}
            if len(headers) != len(set(headers)) or set(headers) != set(expected):
                raise ValueError('Header tidak cocok. Unduh template v2 sesuai akun/PT Anda; template lama tidak diterima.')
            if sheet.max_row and sheet.max_row > 2505:
                raise ValueError('Terlalu banyak baris fisik.')
            rows = []
            for number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers) + 1), 2):
                if number > 2505:
                    raise ValueError('Terlalu banyak baris fisik.')
                if not any(cell.value not in (None, '') for cell in cells):
                    continue
                if cells[-1].value not in (None, '') or any(cell.data_type == 'f' for cell in cells):
                    raise ValueError(f'Baris {number}: formula/kolom tambahan tidak diizinkan.')
                values = {expected[header]: cell.value for header, cell in zip(headers, cells)}
                rows.append((number, values))
                if len(rows) > 500:
                    raise ValueError('Maksimal 500 pegawai per impor.')
        finally:
            workbook.close()
        if not rows:
            raise ValueError('File tidak berisi data pegawai.')
        errors, saved, batches = [], 0, {}
        with changes.write_transaction() as db:
            for number, incoming in rows:
                try:
                    incoming.setdefault('perusahaan_induk', session.get('company') or '')
                    status = str(incoming.get('status_aktif') or 'Aktif').strip().lower()
                    if status not in {'aktif', 'nonaktif', '1', '0'}:
                        raise ValueError('Status harus Aktif/Nonaktif.')
                    incoming['status_aktif'] = int(status in {'aktif', '1'})
                    incoming['siklus_gaji'] = str(incoming.get('siklus_gaji') or 'B').strip().upper()
                    values = changes.validate_employee(db, incoming)
                    changes.assert_identity(db, values)
                    client = changes.scoped_client(db, values['perusahaan_induk'])
                    if session.get('is_superadmin'):
                        changes.create_employee(db, values)
                    else:
                        batch_id = batches.get(client['id'])
                        if not batch_id:
                            batch_id = changes.new_batch(db, client['id'], 'import')
                            batches[client['id']] = batch_id
                        changes.queue_item(db, batch_id, 'create', values)
                        changes.audit(db, client['id'], None, 'requested.import', after=values)
                    saved += 1
                except ValueError as error:
                    errors.append(f'Baris {number}: {error}')
            for client_id, batch_id in batches.items():
                client = dict(db.execute('SELECT * FROM clients WHERE id=?', (client_id,)).fetchone())
                announce(db, client, batch_id, True)
        action = 'ditambahkan' if session.get('is_superadmin') else 'diajukan untuk persetujuan'
        if errors:
            return render_template('shared/employee_import_result.html', saved=saved,
                                   action=action, errors=errors, batches=list(batches.values()))
        flash(f'{saved} pegawai {action}; {len(errors)} baris gagal.', 'success' if saved else 'error')
    except DefusedXmlException:
        flash('File tidak bisa dibaca: XML Excel tidak aman.', 'error')
    except ValueError as error:
        flash(str(error), 'error')
    except (zipfile.BadZipFile, OSError, KeyError, StopIteration):
        flash('File tidak bisa dibaca. Pastikan Excel valid dan tidak rusak.', 'error')
    return redirect(url_for('web.admin_pegawai'))


@bp.post('/admin/pegawai/bulk')
def bulk():
    if not session.get('is_superadmin') and ('admin_fee_flat' in request.form or 'set_admin_fee_flat' in request.form):
        abort(403)
    try:
        with changes.write_transaction() as db:
            token = request.form.get('preview_token')
            if token:
                preview = db.execute('SELECT * FROM employee_bulk_previews WHERE token=? AND actor=?',
                                     (token, changes.actor())).fetchone()
                if not preview or preview['consumed'] or preview['created_at'] < int(time.time()) - 900:
                    raise ValueError('Pratinjau kedaluwarsa/sudah digunakan. Ulangi pratinjau.')
                client = dict(db.execute('SELECT * FROM clients WHERE id=?', (preview['client_id'],)).fetchone())
                changes.scoped_client(db, client['name'])
                targets = json.loads(preview['payload'])
                for target in targets:
                    row = db.execute('SELECT * FROM pegawai WHERE id=?', (target['id'],)).fetchone()
                    if not row or changes.snapshot(row) != target['before']:
                        raise ValueError('Data/target berubah sejak pratinjau. Ulangi pratinjau.')
                batch_id = changes.new_batch(db, client['id'], 'bulk')
                pending = False
                for target in targets:
                    row = db.execute('SELECT * FROM pegawai WHERE id=?', (target['id'],)).fetchone()
                    values = changes.validate_employee(db, target['patch'], row)
                    immediate, queued = changes.submit_update(db, client, row['id'], row, values, batch_id)
                    pending = pending or queued
                db.execute('UPDATE employee_bulk_previews SET consumed=1 WHERE token=?', (token,))
                if not session.get('is_superadmin'):
                    announce(db, client, batch_id, pending)
                flash('Batch tersimpan' + ('; menunggu approval.' if pending else '.'), 'success')
                return redirect(url_for('employees.detail', batch_id=batch_id))
            client = changes.scoped_client(db, request.form.get('client') or session.get('company'))
            patch = {}
            for field in ('gaji', 'plafon', 'admin_fee_flat'):
                if request.form.get('set_' + field) == '1':
                    patch[field] = request.form.get(field, '')
            if 'plafon' in patch:
                patch['plafon_mode'] = request.form.get('plafon_mode')
                patch['plafon_persen'] = request.form.get('plafon_persen')
            if not session.get('is_superadmin') and 'admin_fee_flat' in patch:
                abort(403)
            if not patch:
                raise ValueError('Pilih field yang akan diubah.')
            sql = 'SELECT * FROM pegawai WHERE LOWER(TRIM(perusahaan_induk))=?'
            params = [client['name_key']]
            if request.form.get('all_filtered') == '1':
                query = (request.form.get('q') or '').strip()[:200]
                project = request.form.get('company') or ''
                if query:
                    searchable = ('id_pegawai', 'nama', 'email', 'jabatan', 'perusahaan',
                                  'no_rekening', 'no_rekening_lain', 'rekening_ewallet')
                    sql += ' AND (' + ' OR '.join(f'{field} LIKE ?' for field in searchable) + ')'
                    params += ['%' + query + '%'] * len(searchable)
                if project:
                    sql += ' AND LOWER(TRIM(perusahaan))=?'
                    params.append(name_key(project))
            else:
                ids = sorted({int(value) for value in request.form.getlist('employee_ids')})
                if not ids:
                    raise ValueError('Pilih pegawai dahulu.')
                sql += ' AND id IN (' + ','.join('?' for value in ids) + ')'
                params += ids
            rows = db.execute(sql + ' ORDER BY id', params).fetchall()
            if not rows or (request.form.get('all_filtered') != '1' and len(rows) != len(ids)):
                raise ValueError('Target kosong atau memuat pegawai di luar PT.')
            targets = []
            for row in rows:
                values = changes.validate_employee(db, patch, row)
                actual = {field: values[field] for field in values if changes.snapshot(row)[field] != values[field]}
                if actual:
                    targets.append({'id': row['id'], 'before': changes.snapshot(row), 'patch': actual})
            if not targets:
                raise ValueError('Tidak ada nilai yang berubah.')
            token = secrets.token_urlsafe(32)
            db.execute('''INSERT INTO employee_bulk_previews (token,actor,client_id,payload,created_at)
                VALUES (?,?,?,?,?)''', (token, changes.actor(), client['id'], changes.encoded(targets), int(time.time())))
        return render_template('shared/employee_bulk_preview.html', targets=targets, client=client, token=token)
    except ValueError as error:
        flash(str(error), 'error')
    return redirect(url_for('web.admin_pegawai'))
