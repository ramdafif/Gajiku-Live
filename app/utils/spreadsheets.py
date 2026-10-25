import csv
import io
import math
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SPREADSHEET_MIMETYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def safe_spreadsheet_value(value):
    if isinstance(value, str) and value.lstrip(" ").startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + value
    return value


class SafeCSVWriter:
    def __init__(self, stream, **kwargs):
        self.writer = csv.writer(stream, **kwargs)

    def writerow(self, values):
        return self.writer.writerow(safe_spreadsheet_value(value) for value in values)

    def writerows(self, rows):
        for row in rows:
            self.writerow(row)


def _excel_value(header, value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if header in ("ID Pegawai", "No_Rek Bank") and value is not None:
        return str(value)
    if header not in ("Tanggal", "Dibuat") or not value:
        return value
    if isinstance(value, (date, datetime)):
        parsed = value
    else:
        text = str(value).strip()
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            try:
                pattern = "%d/%m/%Y" if header == "Tanggal" else "%d/%m/%Y %H:%M:%S"
                parsed = datetime.strptime(text, pattern)
            except ValueError:
                return value
    if isinstance(parsed, datetime):
        return parsed.date() if header == "Tanggal" else parsed.replace(tzinfo=None)
    return parsed


def build_spreadsheet_bytes(headers, rows, file_format="xlsx"):
    if file_format == "csv":
        stream = io.StringIO()
        writer = SafeCSVWriter(stream)
        writer.writerow(headers)
        writer.writerows(rows)
        return stream.getvalue().encode("utf-8-sig")
    if file_format != "xlsx":
        raise ValueError("Unsupported spreadsheet format")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transaksi"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(list(headers))
    widths = [len(header) + 4 for header in headers]
    body_font = Font(name="Calibri", size=11)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values):
            header = headers[column_index]
            cell = sheet.cell(row=row_index, column=column_index + 1,
                              value=_excel_value(header, value))
            cell.font = body_font
            cell.alignment = body_alignment
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if header in ("ID Pegawai", "No_Rek Bank"):
                cell.number_format = "@"
            elif header in ("Nominal", "Admin"):
                cell.number_format = '"Rp" #,##0'
            elif header == "Tanggal":
                cell.number_format = "dd/mm/yyyy"
            elif header == "Dibuat":
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            if isinstance(cell.value, datetime):
                display = cell.value.strftime("%d/%m/%Y %H:%M:%S")
            elif isinstance(cell.value, date):
                display = cell.value.strftime("%d/%m/%Y")
            elif header in ("Nominal", "Admin") and cell.value is not None:
                display = f"Rp {cell.value:,.0f}"
            else:
                display = str(cell.value) if cell.value is not None else ""
            content_width = max((len(line) for line in display.splitlines()), default=0)
            widths[column_index] = min(60, max(widths[column_index], content_width + 4))

    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for cells in sheet.iter_rows(min_row=2):
        line_count = max(
            sum(max(1, math.ceil(len(line) / (widths[index] - 4)))
                for line in str(cell.value or "").split("\n"))
            for index, cell in enumerate(cells)
        )
        sheet.row_dimensions[cells[0].row].height = min(409, max(22, line_count * 15 + 6))
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155E63")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    if sheet.max_row > 1:
        table = Table(displayName="TransaksiExport", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
